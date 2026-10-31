"""Synthetic-only trial-balance controls; no canonical finance workbook is opened."""
from __future__ import annotations

from io import BytesIO
import json

import pytest
from openpyxl import Workbook, load_workbook

from aram_backend.coverage_preview import CoverageError
from aram_backend.pilot_workbook import SHEET_HEADERS
from aram_backend.trial_balance_diagnostic import analyze_trial_balance


ROWS = (
    ("CASH", "100", "40", "20", "120"),
    ("PAY", "-80", "10", "20", "-90"),
    ("EQT", "-20", "0", "10", "-30"),
)


def workbook() -> bytes:
    book = Workbook()
    book.remove(book.active)
    for name, headers in SHEET_HEADERS.items():
        book.create_sheet(name).append(list(headers))

    def add(name, **columns):
        book[name].append([columns.get(k) for k in SHEET_HEADERS[name]])

    add("Company", company_code="001", company_name="Synthetic only",
        fiscal_year=1404, status="Draft")
    for i, (code, opening, debit, credit, closing) in enumerate(ROWS, 2):
        add("Accounts", company_code="001", fiscal_year=1404,
            general_code=code, general_name="Synthetic GL",
            account_full_code=code, account_name="Synthetic " + code,
            opening=opening, debit=debit, credit=credit, closing=closing,
            source_row=i + 10, closing_delta="0", row_status="Draft")
    out = BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def edit(data, change):
    book = load_workbook(BytesIO(data))
    change(book)
    result = BytesIO()
    book.save(result)
    book.close()
    return result.getvalue()


def audit(data=None, **scope):
    return analyze_trial_balance(
        workbook() if data is None else data,
        company_code=scope.get("company_code", "001"),
        fiscal_year=scope.get("fiscal_year", 1404),
    )


def test_balanced_opening_movement_closing_without_financial_approval():
    result = audit()
    assert result.account_count == 3
    assert result.opening_debit_net_total == "0"
    assert result.debit_turnover_total == "50"
    assert result.credit_turnover_total == "50"
    assert result.closing_debit_net_total == "0"
    assert result.movement_difference == "0"
    assert result.opening_to_closing_difference == "0"
    assert result.balanced_for_analyst_review
    assert not result.findings
    assert [a.account_code for a in result.accounts] == ["CASH", "EQT", "PAY"]
    assert not result.source_ledger_completeness_verified
    assert not result.accounting_policies_verified
    assert not result.official_financial_statement
    assert not result.independently_reviewed
    assert not result.releasable


def test_unbalanced_opening_and_closing_are_both_reported():
    def change(book):
        book["Accounts"]["Q2"] = "101"
        book["Accounts"]["T2"] = "121"
    result = audit(edit(workbook(), change))
    assert result.opening_debit_net_total == "1"
    assert result.closing_debit_net_total == "1"
    assert result.movement_difference == "0"
    assert not result.balanced_for_analyst_review
    assert result.findings == (
        "OPENING_DEBIT_NET_UNBALANCED", "CLOSING_DEBIT_NET_UNBALANCED"
    )


def test_mismatched_debit_credit_turnover_is_reported():
    def change(book):
        book["Accounts"]["R2"] = "41"
        book["Accounts"]["T2"] = "121"
    result = audit(edit(workbook(), change))
    assert result.debit_turnover_total == "51"
    assert result.credit_turnover_total == "50"
    assert result.movement_difference == "1"
    assert result.closing_debit_net_total == "1"
    assert "DEBIT_CREDIT_TURNOVER_UNBALANCED" in result.findings
    assert "CLOSING_DEBIT_NET_UNBALANCED" in result.findings


def test_inconsistent_account_rollforward_rejected_before_aggregation():
    source = edit(workbook(), lambda book: setattr(book["Accounts"]["T2"], "value", "999"))
    with pytest.raises(CoverageError) as err:
        audit(source)
    assert err.value.code == "TRIAL_BALANCE_WORKBOOK_INVALID"


def test_negative_turnover_rejected_not_net_off():
    def change(book):
        book["Accounts"]["R2"] = "-1"
        book["Accounts"]["T2"] = "79"
    with pytest.raises(CoverageError) as err:
        audit(edit(workbook(), change))
    assert err.value.code == "TRIAL_BALANCE_NEGATIVE_TURNOVER"


def test_all_zero_rows_are_valid_zero_math_without_claiming_completeness():
    def change(book):
        for row in book["Accounts"].iter_rows(min_row=2):
            for index in (16, 17, 18, 19):
                row[index].value = "0"
    result = audit(edit(workbook(), change))
    assert result.balanced_for_analyst_review
    assert result.opening_debit_net_total == "0"
    assert result.debit_turnover_total == "0"
    assert result.credit_turnover_total == "0"
    assert result.closing_debit_net_total == "0"
    assert not result.source_ledger_completeness_verified


def test_scope_excludes_other_company_even_if_other_is_unbalanced():
    def change(book):
        headers = SHEET_HEADERS["Company"]
        book["Company"].append([{
            "company_code": "002", "company_name": "Other",
            "fiscal_year": 1404, "status": "Draft",
        }.get(key) for key in headers])
        book["Accounts"].append([{
            "company_code": "002", "fiscal_year": 1404,
            "general_code": "A", "general_name": "Other",
            "account_full_code": "A", "account_name": "Other",
            "opening": "999", "debit": "0", "credit": "0",
            "closing": "999", "source_row": 14,
            "closing_delta": "0", "row_status": "Draft",
        }.get(key) for key in SHEET_HEADERS["Accounts"]])
    data = edit(workbook(), change)
    assert audit(data).balanced_for_analyst_review
    other = audit(data, company_code="002")
    assert not other.balanced_for_analyst_review
    assert other.account_count == 1


def test_duplicate_account_fails_closed():
    def change(book):
        sheet = book["Accounts"]
        sheet.append([c.value for c in sheet[2]])
    with pytest.raises(CoverageError):
        audit(edit(workbook(), change))


def test_hidden_auxiliary_sheet_fails_closed():
    data = edit(workbook(), lambda book: setattr(book["Checks"], "sheet_state", "hidden"))
    with pytest.raises(CoverageError) as err:
        audit(data)
    assert err.value.code == "HIDDEN_INPUT"


def test_hidden_account_row_fails_closed():
    data = edit(workbook(), lambda book: setattr(book["Accounts"].row_dimensions[2], "hidden", True))
    with pytest.raises(CoverageError) as err:
        audit(data)
    assert err.value.code == "HIDDEN_INPUT"


def test_merged_accounts_fail_closed():
    data = edit(workbook(), lambda book: book["Accounts"].merge_cells("C2:D2"))
    with pytest.raises(CoverageError) as err:
        audit(data)
    assert err.value.code in {"MERGED_INPUT", "TRIAL_BALANCE_WORKBOOK_INVALID"}


def test_wrong_header_rejected():
    data = edit(workbook(), lambda book: setattr(book["Accounts"]["A1"], "value", "other"))
    with pytest.raises(CoverageError) as err:
        audit(data)
    assert err.value.code == "TRIAL_BALANCE_WORKBOOK_INVALID"


def test_scope_must_exist_no_fallback():
    with pytest.raises(CoverageError) as err:
        audit(company_code="999")
    assert err.value.code == "TRIAL_BALANCE_COMPANY_NOT_FOUND"
    with pytest.raises(CoverageError):
        audit(fiscal_year=1403)


@pytest.mark.parametrize("bad_company,bad_year", [
    ("", 1404), ("001", 0), ("001", True),
    ("001", "1404"), (None, 1404),
])
def test_invalid_scope_fails_closed(bad_company, bad_year):
    with pytest.raises(CoverageError):
        audit(company_code=bad_company, fiscal_year=bad_year)


def test_reject_unsafe_non_xlsx_inputs_and_oversize():
    with pytest.raises(CoverageError):
        audit(b"invalid")
    with pytest.raises(CoverageError):
        audit(b"x" * (20 * 1024 * 1024 + 1))
    with pytest.raises(CoverageError):
        audit("not-bytes")


def test_scope_with_no_accounts_rejected():
    def change(book):
        book["Accounts"].delete_rows(2, 3)
    with pytest.raises(CoverageError) as err:
        audit(edit(workbook(), change))
    assert err.value.code == "TRIAL_BALANCE_NO_ACCOUNTS"


def test_json_safe_values_and_input_bytes_unchanged():
    original = workbook()
    snapshot = bytes(original)
    result = json.loads(json.dumps(audit(original).to_dict()))
    assert original == snapshot
    assert result["releasable"] is False
    assert len(result["source_sha256"]) == 64
    assert result["accounts"][0]["opening"] == "100"


def test_account_rows_include_source_lineage_and_net_movement():
    result = audit()
    cash = next(a for a in result.accounts if a.account_code == "CASH")
    assert cash.excel_row == 2
    assert cash.source_row == 12
    assert cash.debit_net_movement == "20"
    assert cash.closing == "120"
