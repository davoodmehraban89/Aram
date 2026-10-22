from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from aram_backend.mapping_excel import (
    MappingEntry, export_mapping_workbook, import_mapping_workbook, preview_mapping_workbook,
)
from aram_backend.mapping_service import MappingValidationError
from aram_backend.models import AccountVersion, Company, FiscalYear, MappingRule, NoteDefinition


def make_year(session, code):
    company = Company(code=code, name="Test company")
    year = FiscalYear(company=company, year=1404)
    session.add(year)
    session.flush()
    session.add(AccountVersion(year_id=year.id, account_code="001", account_name="Cash", version=1))
    note = NoteDefinition(year_id=year.id, note_number="1", title="Cash", version=1)
    session.add(note)
    session.commit()
    return year, note


def entry(note_id, **overrides):
    values = dict(rule_key="cash", account_code="001", target_note_id=note_id,
                  mapping_type="direct", version=1)
    values.update(overrides)
    return MappingEntry(**values)


def xlsx(rows, headers=("rule_key", "account_code", "target_note_id", "mapping_type", "version")):
    book = Workbook()
    book.active.title = "Mapping"
    book.active.append(headers)
    for row in rows:
        book.active.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_round_trip_and_commit(session):
    year, note = make_year(session, "EXCEL-A")
    workbook = export_mapping_workbook((entry(note.id),))
    assert preview_mapping_workbook(workbook) == (entry(note.id),)
    assert import_mapping_workbook(session, year_id=year.id, content=workbook) == 1
    session.commit()
    assert session.scalar(select(MappingRule)).account_code == "001"


def test_reimport_is_idempotent(session):
    year, note = make_year(session, "EXCEL-B")
    workbook = export_mapping_workbook((entry(note.id),))
    import_mapping_workbook(session, year_id=year.id, content=workbook)
    import_mapping_workbook(session, year_id=year.id, content=workbook)
    session.commit()
    assert len(session.scalars(select(MappingRule)).all()) == 1


def test_rollback_when_later_row_invalid(session):
    year, note = make_year(session, "EXCEL-C")
    workbook = export_mapping_workbook((entry(note.id), entry(note.id, rule_key="missing", account_code="999")))
    with pytest.raises(MappingValidationError, match="account code"):
        import_mapping_workbook(session, year_id=year.id, content=workbook)
    assert session.scalars(select(MappingRule)).all() == []


def test_cross_year_note_rolls_back(session):
    year, note = make_year(session, "EXCEL-D")
    _, other_note = make_year(session, "EXCEL-E")
    workbook = export_mapping_workbook((entry(note.id), entry(other_note.id, rule_key="foreign")))
    with pytest.raises(MappingValidationError, match="same fiscal year"):
        import_mapping_workbook(session, year_id=year.id, content=workbook)
    assert session.scalars(select(MappingRule)).all() == []


def test_caller_can_rollback_successful_import(session):
    year, note = make_year(session, "EXCEL-F")
    import_mapping_workbook(session, year_id=year.id, content=export_mapping_workbook((entry(note.id),)))
    session.rollback()
    assert session.scalars(select(MappingRule)).all() == []


def test_formula_injection_is_not_executable_in_export():
    workbook = export_mapping_workbook((entry(1, rule_key="=2+2", mapping_type="@HYPERLINK"),))
    sheet = load_workbook(BytesIO(workbook)).active
    assert sheet["A2"].data_type == "s"
    assert sheet["D2"].data_type == "s"
    assert preview_mapping_workbook(workbook)[0].rule_key == "=2+2"


@pytest.mark.parametrize("rows,pattern", [
    ([("cash", "001", 1, "direct", 1), ("cash", "001", 1, "direct", 1)], "duplicate"),
    ([("cash", "001", "1", "direct", 1)], "positive integer"),
    ([("cash", "001", True, "direct", 1)], "positive integer"),
    ([("cash", "001", 1, "direct", 0)], "positive integer"),
    ([(" cash", "001", 1, "direct", 1)], "invalid rule_key"),
    ([("cash", "001", 1, "direct", "=1+1")], "formulas"),
    ([], "no mapping entries"),
])
def test_reject_bad_workbook_rows(rows, pattern):
    with pytest.raises(MappingValidationError, match=pattern):
        preview_mapping_workbook(xlsx(rows))


def test_reject_wrong_headers():
    with pytest.raises(MappingValidationError, match="headers"):
        preview_mapping_workbook(xlsx([("cash", "001", 1, "direct", 1)], headers=("wrong",)))


def test_reject_additional_sheet():
    book = Workbook()
    book.active.title = "Mapping"
    book.active.append(("rule_key", "account_code", "target_note_id", "mapping_type", "version"))
    book.create_sheet("Hidden")
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(MappingValidationError, match="only a Mapping sheet"):
        preview_mapping_workbook(stream.getvalue())


@pytest.mark.parametrize("content", [b"", b"not an xlsx", b"x" * (10 * 1024 * 1024 + 1)])
def test_reject_invalid_archives(content):
    with pytest.raises(MappingValidationError):
        preview_mapping_workbook(content)


def test_reject_decompression_bomb():
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("huge.txt", b"x" * (40 * 1024 * 1024 + 1))
    with pytest.raises(MappingValidationError, match="safety limits"):
        preview_mapping_workbook(stream.getvalue())


def test_reject_duplicate_export():
    with pytest.raises(MappingValidationError, match="duplicate"):
        export_mapping_workbook((entry(1), entry(1)))


def test_reject_empty_export():
    with pytest.raises(MappingValidationError, match="1 to 10000"):
        export_mapping_workbook(())
