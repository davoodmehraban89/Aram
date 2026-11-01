from dataclasses import replace
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from openpyxl import load_workbook

from aram_backend.mapping_store import MappingStore
from aram_backend.mapping_workflow import (
    MappingDecision, MappingWorkflowError, export_pilot_mapping_workbook,
    import_mapping_workbook,
)
from test_pilot_workbook import valid_rows, workbook_bytes


def decision():
    return MappingDecision("M001", "001", 1405, "0010", "BS", "Assets",
                           "01", "001", 1, "Draft", None, "Manual")


def test_durable_reopen_audit_and_optimistic_concurrency(tmp_path):
    db = tmp_path / "mapping.sqlite3"
    store = MappingStore(db)
    first, version = store.load("001", 1405)
    second, other_version = store.load("001", 1405)
    assert version == other_version == 0
    first.propose(decision(), actor="maker")
    assert store.save(first, version) == 1
    with pytest.raises(MappingWorkflowError, match="stale"):
        store.save(second, other_version)
    reopened = MappingStore(db)
    state, version = reopened.load("001", 1405)
    assert version == 1 and state.decisions[0] == decision()
    state.transition("M001", "Reviewed", actor="reviewer")
    assert reopened.save(state, version) == 2
    assert len(reopened.audit_log("001", 1405)) == 2
    with pytest.raises(MappingWorkflowError, match="stale"):
        store.save(first, 1)


def test_two_writers_one_wins(tmp_path):
    store = MappingStore(tmp_path / "concurrent.db")
    first, version = store.load("001", 1405)
    first.propose(decision(), actor="maker")
    assert store.save(first, version) == 1
    def worker(actor):
        state, version = store.load("001", 1405)
        state.transition("M001", "Reviewed", actor=actor)
        return state, version
    a = worker("reviewer-a")
    b = worker("reviewer-b")
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(store.save, state, revision) for state, revision in (a, b)]
        outcomes = []
        for future in futures:
            try:
                outcomes.append(future.result())
            except MappingWorkflowError:
                outcomes.append("conflict")
    assert sorted(map(str, outcomes)) == ["2", "conflict"]
    assert len(store.audit_log("001", 1405)) == 2


def test_append_only_audit_and_scope(tmp_path):
    store = MappingStore(tmp_path / "audit.db")
    space, rev = store.load("001", 1405)
    space.propose(decision(), actor="a")
    store.save(space, rev)
    loaded, rev = store.load("001", 1405)
    loaded._audit.clear()
    with pytest.raises(MappingWorkflowError, match="append-only"):
        store.save(loaded, rev)
    assert len(store.audit_log("001", 1405)) == 1
    other, rev = store.load("002", 1405)
    assert rev == 0 and other.decisions == ()


def test_full_workbook_round_trip_preserves_other_sheets():
    original = workbook_bytes(valid_rows())
    imported = import_mapping_workbook(original, "001", 1405)
    exported = export_pilot_mapping_workbook(original, imported, company_code="001", fiscal_year=1405)
    assert import_mapping_workbook(exported, "001", 1405) == imported
    before = load_workbook(BytesIO(original), read_only=True)
    after = load_workbook(BytesIO(exported), read_only=True)
    try:
        for name in before.sheetnames:
            if name == "Mapping":
                continue
            assert list(before[name].values) == list(after[name].values)
    finally:
        before.close()
        after.close()


def test_export_rejects_foreign_and_unverifiable_approval():
    original = workbook_bytes(valid_rows())
    with pytest.raises(MappingWorkflowError, match="foreign"):
        export_pilot_mapping_workbook(original, [replace(decision(), company_code="002")],
                                      company_code="001", fiscal_year=1405)
    with pytest.raises(MappingWorkflowError, match="review"):
        export_pilot_mapping_workbook(original, [replace(decision(), status="Approved")],
                                      company_code="001", fiscal_year=1405)
