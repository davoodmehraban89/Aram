from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from aram_backend.pilot_workbook import SHEET_HEADERS, preflight_pilot_workbook


def workbook_bytes(rows: dict[str, list[tuple]] | None = None) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, headers in SHEET_HEADERS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in (rows or {}).get(name, []):
            sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def valid_rows() -> dict[str, list[tuple]]:
    company = ("001", "آرام", 1405, "IRR", "rial", "tester", "Draft")
    account = (
        "001", 1405, "10", "کل", None, None, None, None, None, None, None,
        None, None, None, "0010", "حساب", 10, 5, 3, 12, 2, "Draft", 0,
    )
    note = ("001", 1405, "BS", "Assets", "01", "یادداشت", "001", "زیر یادداشت")
    mapping = (
        "M001", "001", 1405, "0010", "General", "10", None, None, None,
        None, None, "BS", "Assets", "01", "001", 1, "Draft", None, None,
        "Manual", None, "Pending",
    )
    return {"Company": [company], "Accounts": [account], "Notes": [note], "Mapping": [mapping]}


def issue_codes(data: bytes) -> list[str]:
    return [issue.code for issue in preflight_pilot_workbook(data).issues]


def test_minimal_valid_workbook_preserves_input_bytes():
    data = workbook_bytes(valid_rows())
    original = bytes(data)
    result = preflight_pilot_workbook(data)
    assert result.valid
    assert result.issues == ()
    assert data == original


def test_invalid_workbook_bytes_return_a_typed_error():
    result = preflight_pilot_workbook(b"not an xlsx")
    assert not result.valid
    assert result.issues[0].code == "INVALID_WORKBOOK"


def test_committed_blank_template_is_valid():
    template = Path(__file__).parents[2] / "samples" / "Aram_Pilot_Input_and_Mapping_Template_v0.1.xlsx"
    result = preflight_pilot_workbook(template.read_bytes())
    assert result.valid, result.issues


def test_missing_sheet_and_wrong_header_order_report_exact_location():
    data = workbook_bytes()
    book = load_workbook(BytesIO(data))
    del book["Notes"]
    book["Company"].move_range("A1:B1", cols=1)
    stream = BytesIO()
    book.save(stream)
    result = preflight_pilot_workbook(stream.getvalue())
    assert ("HEADER_MISMATCH", "Company", 1, "__headers__") in {
        (i.code, i.sheet, i.row, i.field) for i in result.issues
    }
    assert ("MISSING_SHEET", "Notes", 1, "__sheet__") in {
        (i.code, i.sheet, i.row, i.field) for i in result.issues
    }


def test_numeric_code_rejected_and_leading_zero_text_accepted():
    rows = valid_rows()
    account = list(rows["Accounts"][0])
    account[14] = 10
    rows["Accounts"] = [tuple(account)]
    result = preflight_pilot_workbook(workbook_bytes(rows))
    assert ("CODE_NOT_TEXT", "account_full_code") in {(i.code, i.field) for i in result.issues}

    valid = preflight_pilot_workbook(workbook_bytes(valid_rows()))
    assert valid.valid


def test_required_duplicate_and_cross_reference_errors_are_collected():
    rows = valid_rows()
    rows["Company"].append(rows["Company"][0])
    rows["Accounts"].append(rows["Accounts"][0])
    broken_note = list(rows["Notes"][0])
    broken_note[0] = "999"
    broken_note[5] = None
    rows["Notes"] = [tuple(broken_note)]
    codes = issue_codes(workbook_bytes(rows))
    assert "DUPLICATE_COMPANY_YEAR" in codes
    assert "DUPLICATE_ACCOUNT" in codes
    assert "UNKNOWN_COMPANY_YEAR" in codes
    assert "REQUIRED_FIELD" in codes


def test_unknown_account_duplicate_mapping_and_lifecycle_errors():
    rows = valid_rows()
    mapping = list(rows["Mapping"][0])
    mapping[3] = "0099"
    mapping[16] = "Published"
    rows["Mapping"] = [tuple(mapping), tuple(mapping)]
    codes = issue_codes(workbook_bytes(rows))
    assert "UNKNOWN_MAPPING_ACCOUNT" in codes
    assert "DUPLICATE_MAPPING_ID" in codes
    assert "INVALID_LIFECYCLE" in codes


def test_non_finite_values_and_non_zero_closing_delta():
    rows = valid_rows()
    account = list(rows["Accounts"][0])
    account[16] = "NaN"
    account[19] = 99
    rows["Accounts"] = [tuple(account)]
    codes = issue_codes(workbook_bytes(rows))
    assert "NON_FINITE_FINANCIAL_VALUE" in codes

    rows = valid_rows()
    account = list(rows["Accounts"][0])
    account[19] = 11
    rows["Accounts"] = [tuple(account)]
    assert "NON_ZERO_CLOSING_DELTA" in issue_codes(workbook_bytes(rows))


def test_positive_integer_fields_and_deterministic_ordering():
    rows = valid_rows()
    account = list(rows["Accounts"][0])
    account[20] = 1.5
    rows["Accounts"] = [tuple(account)]
    mapping = list(rows["Mapping"][0])
    mapping[15] = 0
    rows["Mapping"] = [tuple(mapping)]
    data = workbook_bytes(rows)
    first = preflight_pilot_workbook(data)
    second = preflight_pilot_workbook(data)
    assert first.issues == second.issues
    assert [i.code for i in first.issues].count("INVALID_POSITIVE_INTEGER") == 2
    assert first.issue_counts == {"ERROR": len(first.issues)}


def test_validation_does_not_persist_any_artifact(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    before = set(tmp_path.iterdir())
    preflight_pilot_workbook(workbook_bytes(valid_rows()))
    assert set(tmp_path.iterdir()) == before
