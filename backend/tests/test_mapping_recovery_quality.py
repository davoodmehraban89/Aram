import sqlite3
from dataclasses import replace

import pytest
from aram_backend.mapping_quality import assess_mapping_quality
from aram_backend.mapping_store import MappingStore
from aram_backend.mapping_workflow import MappingDecision, MappingWorkflowError


def decision(code="0010", mapping_id="M1", status="Draft"):
    return MappingDecision(mapping_id, "001", 1405, code, "BS", "Assets",
                           "01", None, 1, status, None, "Manual")


def test_backup_restore_to_fresh_store_and_no_overwrite(tmp_path):
    store = MappingStore(tmp_path / "original.db")
    state, revision = store.load("001", 1405)
    state.propose(decision(), actor="maker")
    store.save(state, revision)
    target = tmp_path / "backup.db"
    store.backup_to(target)
    restored = MappingStore(target)
    restored.verify_integrity()
    assert restored.load("001", 1405)[0].decisions == (decision(),)
    assert restored.audit_log("001", 1405) == store.audit_log("001", 1405)
    with pytest.raises(MappingWorkflowError, match="new"):
        store.backup_to(target)


def test_detects_direct_audit_tamper(tmp_path):
    path = tmp_path / "tamper.db"
    store = MappingStore(path)
    state, rev = store.load("001", 1405)
    state.propose(decision(), actor="maker")
    store.save(state, rev)
    with sqlite3.connect(path) as db:
        db.execute("DELETE FROM mapping_audit")
    with pytest.raises(MappingWorkflowError, match="audit/state mismatch"):
        store.verify_integrity()


def test_quality_never_auto_approves_ambiguous_accounts():
    draft = decision()
    approved = replace(draft, status="Approved", version=3)
    conflict = replace(approved, mapping_id="M2")
    report = assess_mapping_quality(["0010", "0020"], [approved, conflict],
                                    company_code="001", fiscal_year=1405)
    assert report.conflicting_accounts == ("0010",)
    assert report.unmapped_accounts == ("0020",)
    assert not report.ready
    report = assess_mapping_quality(["0010"], [approved],
                                    company_code="001", fiscal_year=1405)
    assert report.ready and report.approved_accounts == ("0010",)
    report = assess_mapping_quality(["0010"], [draft],
                                    company_code="001", fiscal_year=1405)
    assert report.pending_accounts == ("0010",) and not report.ready
    with pytest.raises(MappingWorkflowError, match="foreign"):
        assess_mapping_quality(["0010"], [replace(draft, company_code="002")],
                               company_code="001", fiscal_year=1405)


def test_workbook_mapping_notes_preserved_on_roundtrip():
    from io import BytesIO

    from aram_backend.mapping_workflow import (
        export_pilot_mapping_workbook,
        import_mapping_workbook,
    )
    from aram_backend.pilot_workbook import SHEET_HEADERS
    from openpyxl import load_workbook
    from test_pilot_workbook import valid_rows, workbook_bytes

    wb = load_workbook(BytesIO(workbook_bytes(valid_rows())))
    columns = SHEET_HEADERS["Mapping"]
    wb["Mapping"].cell(2, columns.index("notes") + 1, "Manual financial review")
    data = BytesIO()
    wb.save(data)
    wb.close()
    decisions = import_mapping_workbook(data.getvalue(), "001", 1405)
    output = export_pilot_mapping_workbook(data.getvalue(), decisions,
                                           company_code="001", fiscal_year=1405)
    result = load_workbook(BytesIO(output), read_only=True)
    try:
        assert result["Mapping"].cell(2, columns.index("notes") + 1).value == "Manual financial review"
    finally:
        result.close()
