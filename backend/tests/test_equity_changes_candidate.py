"""Synthetic-only tests for the unapproved statement of changes in equity."""
from __future__ import annotations

from dataclasses import replace
from io import BytesIO
import json

import pytest
from openpyxl import Workbook, load_workbook

from aram_backend.coverage_preview import CoverageError
from aram_backend.equity_changes_candidate import (
    EquityMovementPolicy, build_equity_changes_candidate,
)
from aram_backend.pilot_workbook import SHEET_HEADERS
from aram_backend.reconciliation_review import NoteReference


ITEMS = (
    ("CAP", "Capital", "C01", "-100", "0", "20", "-120",
     "SHARE_CAPITAL", "CAPITAL_CONTRIBUTION"),
    ("RET", "Retained earnings", "R01", "-50", "0", "5", "-55",
     "RETAINED_EARNINGS", "NET_PROFIT_LOSS"),
    ("DIV", "Dividends", "D01", "0", "20", "0", "20",
     "RETAINED_EARNINGS", "DIVIDENDS"),
)


def workbook(*, include_eq=True, bad_delta=False, bad_mapping=False,
             negative_turnover=False):
    book = Workbook()
    book.remove(book.active)
    for name, headers in SHEET_HEADERS.items():
        book.create_sheet(name).append(list(headers))

    def add(name, **values):
        book[name].append([values.get(key) for key in SHEET_HEADERS[name]])

    add("Company", company_code="001", company_name="Synthetic only",
        fiscal_year=1404, currency="IRR", unit="Rial", status="Draft")
    for index, item in enumerate(ITEMS, 2):
        code, label, note, opening, debit, credit, closing, component, movement = item
        if bad_delta and code == "CAP":
            closing = "-119"
        if negative_turnover and code == "DIV":
            debit = "-20"
        add("Accounts", company_code="001", fiscal_year=1404,
            general_code=code, general_name=label, account_full_code=code,
            account_name="Synthetic " + label, opening=opening,
            debit=debit, credit=credit, closing=closing,
            source_row=index + 10, row_status="Draft",
            closing_delta="0" if not bad_delta else "1")
        if include_eq:
            add("Notes", company_code="001", fiscal_year=1404,
                statement="EQ", statement_section=label,
                note_number=note, note_title=label, subnote_code="",
                subnote_title="")
            if not (bad_mapping and code == "DIV"):
                add("Mapping", mapping_id="M-" + code, company_code="001",
                    fiscal_year=1404, account_full_code=code,
                    account_level="General", general_code=code,
                    target_statement="EQ", statement_section=label,
                    note_number=note, subnote_code="", mapping_version=1,
                    approval_status="Approved", reviewer="Synthetic reviewer",
                    review_date="2026-10-31", source="Synthetic",
                    validation_status="Valid")
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def references():
    return [
        NoteReference("EQ", label, note, "", closing, "SYNTH-REF-" + code)
        for code, label, note, opening, debit, credit, closing, component, movement in ITEMS
    ]


def policies():
    return [
        EquityMovementPolicy(
            label, note, "", component, movement, -1,
            "Synthetic " + label, "SYNTH-POL-" + code,
        )
        for code, label, note, opening, debit, credit, closing, component, movement in ITEMS
    ]


def build(*, data=None, refs=None, pol=None, company="001", year=1404):
    return build_equity_changes_candidate(
        workbook() if data is None else data,
        company_code=company, fiscal_year=year,
        references=references() if refs is None else refs,
        policies=policies() if pol is None else pol,
    )


def edit(data, change):
    book = load_workbook(BytesIO(data))
    change(book)
    out = BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def test_equity_statement_rollforward_and_category_totals_are_exact():
    result = build()
    assert result.eligible_for_analyst_review
    assert not result.findings
    assert result.opening_equity_total == "150"
    assert result.capital_contributions == "20"
    assert result.net_profit_loss_movements == "5"
    assert result.dividends == "-20"
    assert result.other_comprehensive_income == "0"
    assert result.transfers == "0"
    assert result.prior_period_adjustments == "0"
    assert result.other_movements == "0"
    assert result.total_movement == "5"
    assert result.closing_equity_total == "155"
    assert result.rollforward_difference == "0"
    assert len(result.note_movements) == 3
    assert len(result.account_movements) == 3
    assert result.kind == "UNAPPROVED_EQUITY_CHANGES_PRESENTATION_CANDIDATE"
    assert not result.official_financial_statement
    assert not result.accounting_policy_authorized
    assert not result.source_ledger_completeness_verified
    assert not result.independently_reviewed
    assert not result.export_authorized
    assert not result.releasable
    retained = next(x for x in result.component_totals
                    if x.equity_component == "RETAINED_EARNINGS")
    assert retained.opening_balance == "50"
    assert retained.movement_amount == "-15"
    assert retained.closing_balance == "35"


def test_zero_movement_and_zero_balances_are_valid():
    def zero(book):
        for row in book["Accounts"].iter_rows(min_row=2):
            for column in ("opening", "debit", "credit", "closing"):
                row[16 + ("opening", "debit", "credit", "closing").index(column)].value = "0"
    data = edit(workbook(), zero)
    refs = [replace(ref, expected_closing_exposure="0") for ref in references()]
    result = build(data=data, refs=refs)
    assert result.eligible_for_analyst_review
    assert result.opening_equity_total == "0"
    assert result.total_movement == "0"
    assert result.closing_equity_total == "0"


def test_missing_policy_blocks_entire_equity_candidate():
    result = build(pol=policies()[:-1])
    assert not result.eligible_for_analyst_review
    assert "MISSING_EQUITY_POLICY" in result.findings


def test_orphan_policy_blocks_candidate():
    p = policies()
    p.append(EquityMovementPolicy(
        "Unknown", "X01", "", "OTHER_EQUITY", "OTHER_MOVEMENT", -1,
        "Synthetic extra", "SYNTH-EXTRA",
    ))
    result = build(pol=p)
    assert not result.eligible_for_analyst_review
    assert "ORPHAN_EQUITY_POLICY" in result.findings


def test_missing_mapping_blocks_even_when_other_equity_notes_match():
    result = build(data=workbook(bad_mapping=True))
    assert not result.eligible_for_analyst_review
    assert "UPSTREAM_RECONCILIATION_BLOCKED" in result.findings


def test_reference_mismatch_blocks_candidate():
    refs = references()
    refs[0] = replace(refs[0], expected_closing_exposure="-121")
    result = build(refs=refs)
    assert not result.eligible_for_analyst_review
    assert "UPSTREAM_RECONCILIATION_BLOCKED" in result.findings


def test_invalid_source_rollforward_fails_closed():
    with pytest.raises(CoverageError):
        build(data=workbook(bad_delta=True))


def test_negative_turnover_fails_closed():
    with pytest.raises(CoverageError) as error:
        build(data=workbook(negative_turnover=True))
    assert error.value.code == "EQUITY_NEGATIVE_TURNOVER"


@pytest.mark.parametrize("invalid", [
    EquityMovementPolicy("Capital", "C01", "", "BAD", "CAPITAL_CONTRIBUTION",
                         -1, "Capital", "P"),
    EquityMovementPolicy("Capital", "C01", "", "SHARE_CAPITAL", "BAD",
                         -1, "Capital", "P"),
    EquityMovementPolicy("Capital", "C01", "", "SHARE_CAPITAL", "DIVIDENDS",
                         True, "Capital", "P"),
    EquityMovementPolicy("Capital", "C01", "", "SHARE_CAPITAL", "DIVIDENDS",
                         0, "Capital", "P"),
    EquityMovementPolicy("Capital", "C01", "", "SHARE_CAPITAL", "DIVIDENDS",
                         -1, "", "P"),
    EquityMovementPolicy("Capital", "C01", "", "SHARE_CAPITAL", "DIVIDENDS",
                         -1, "Capital", ""),
    EquityMovementPolicy("Capital", "C01", "", "SHARE_CAPITAL", "DIVIDENDS",
                         -1, "Bad" + chr(10) + "label", "P"),
])
def test_invalid_equity_policy_fails_closed(invalid):
    with pytest.raises(CoverageError):
        build(pol=[invalid])


def test_duplicate_policy_fails_closed():
    p = policies()
    with pytest.raises(CoverageError) as error:
        build(pol=p + [p[0]])
    assert error.value.code == "DUPLICATE_EQUITY_POLICY"


def test_wrong_company_and_year_are_not_guessed():
    with pytest.raises(CoverageError):
        build(company="999")
    with pytest.raises(CoverageError):
        build(year=1403)


def test_missing_equity_notes_blocks():
    data = workbook(include_eq=False)
    result = build(data=data, refs=[], pol=[])
    assert not result.eligible_for_analyst_review
    assert "NO_EQUITY_NOTES" in result.findings
    assert result.opening_equity_total == "0"
    assert result.closing_equity_total == "0"


def test_source_unchanged_and_json_safe():
    data = workbook()
    original = bytes(data)
    result = json.loads(json.dumps(build(data=data).to_dict()))
    assert data == original
    assert result["releasable"] is False
    assert len(result["source_sha256"]) == 64
    assert len(result["note_references_sha256"]) == 64
    assert len(result["presentation_policies_sha256"]) == 64
    assert all(isinstance(item["movement_amount"], str)
               for item in result["note_movements"])


def test_policy_order_does_not_change_hash_or_result():
    p = policies()
    first = build(pol=p)
    second = build(pol=list(reversed(p)))
    assert first.note_movements == second.note_movements
    assert first.presentation_policies_sha256 == second.presentation_policies_sha256


def test_policy_change_changes_policy_hash():
    p = policies()
    first = build(pol=p)
    p[0] = replace(p[0], line_label="Revised synthetic capital")
    second = build(pol=p)
    assert first.presentation_policies_sha256 != second.presentation_policies_sha256


def test_wrong_source_hash_cannot_be_silently_reused():
    # References are tied to note keys/amounts, while the source snapshot hash
    # is recalculated for each build; a source mutation still needs exact
    # reconciliation, so mutate a mapped account's closing without its control.
    data = edit(workbook(), lambda book: setattr(book["Accounts"]["T2"], "value", "-121"))
    with pytest.raises(CoverageError):
        build(data=data)
