"""Read-only Mapping coverage preview, using synthetic Pilot XLSX inputs only."""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook

from aram_backend.coverage_preview import CoverageError, preview_mapping_coverage
from aram_backend.pilot_workbook import SHEET_HEADERS


def sample_workbook(*, extra_mapping=False, approved=True):
    book = Workbook()
    book.remove(book.active)
    sheets = {}
    for name, headers in SHEET_HEADERS.items():
        sheet = book.create_sheet(name)
        sheet.append(list(headers))
        sheets[name] = sheet

    def add(sheet_name, **fields):
        headers = SHEET_HEADERS[sheet_name]
        sheets[sheet_name].append([fields.get(header) for header in headers])

    add("Company", company_code="001", company_name="Synthetic Company",
        fiscal_year=1404, status="Draft")
    for code, amount in (("0010", "110.25"), ("0020", "15.75")):
        add("Accounts", company_code="001", fiscal_year=1404,
            general_code=code, general_name="General",
            account_full_code=code, account_name="Synthetic " + code,
            opening=amount, debit="0", credit="0", closing=amount,
            source_row=10, row_status="Draft", closing_delta="0")
    add("Notes", company_code="001", fiscal_year=1404, statement="BS",
        statement_section="Assets", note_number="N01", note_title="Cash")
    options = {
        "company_code": "001", "fiscal_year": 1404,
        "account_full_code": "0010", "account_level": "General",
        "general_code": "0010", "target_statement": "BS",
        "statement_section": "Assets", "note_number": "N01",
        "mapping_version": 1, "approval_status": "Approved" if approved else "Draft",
        "source": "Manual", "validation_status": "Valid",
        "reviewer": "Synthetic reviewer" if approved else None,
        "review_date": "2026-10-31" if approved else None,
    }
    add("Mapping", mapping_id="M001", **options)
    if extra_mapping:
        add("Mapping", mapping_id="M002", **options)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def make_change(content, change):
    book = __import__("openpyxl").load_workbook(BytesIO(content))
    change(book)
    result = BytesIO()
    book.save(result)
    book.close()
    return result.getvalue()


def preview(content):
    return preview_mapping_coverage(content, company_code="001", fiscal_year=1404)


def test_includes_only_file_marked_approved_unique_mapping():
    result = preview(sample_workbook())
    assert result.account_count == 2
    assert result.candidate_mapped_accounts == 1
    assert result.blocked_accounts == 1
    assert result.source_closing_exposure == "126.00"
    assert result.candidate_closing_exposure == "110.25"
    assert result.blocked_closing_exposure == "15.75"
    assert len(result.notes) == 1
    assert result.notes[0].note_number == "N01"
    assert result.notes[0].account_codes == ("0010",)
    assert [(issue.code, issue.account_code) for issue in result.exceptions] == [
        ("UNMAPPED_ACCOUNT", "0020")
    ]
    assert not result.financial_statement
    assert not result.trusted_approvals
    assert not result.releasable


def test_multiple_mapping_rows_are_blocked_not_silently_deduplicated():
    result = preview(sample_workbook(extra_mapping=True))
    assert result.candidate_mapped_accounts == 0
    assert result.blocked_accounts == 2
    assert {issue.code for issue in result.exceptions} == {
        "AMBIGUOUS_ACCOUNT_MAPPING", "UNMAPPED_ACCOUNT"
    }
    match = next(e for e in result.exceptions if e.code == "AMBIGUOUS_ACCOUNT_MAPPING")
    assert match.mapping_excel_rows == (2, 3)


def test_unapproved_file_label_is_not_evidence_of_authorization():
    result = preview(sample_workbook(approved=False))
    assert result.candidate_mapped_accounts == 0
    assert result.notes == ()
    assert any(e.code == "MAPPING_NOT_FILE_MARKED_APPROVED" for e in result.exceptions)


@pytest.mark.parametrize("kwargs", [
    {"company_code": "", "fiscal_year": 1404},
    {"company_code": "001", "fiscal_year": 0},
    {"company_code": "001", "fiscal_year": True},
    {"company_code": "001", "fiscal_year": 1403},
])
def test_scope_is_explicit_and_no_fallback(kwargs):
    with pytest.raises(CoverageError):
        preview_mapping_coverage(sample_workbook(), **kwargs)


def test_malformed_workbook_fails_closed():
    with pytest.raises(CoverageError) as exc:
        preview(b"invalid file")
    assert exc.value.code in {"WORKBOOK_INVALID", "WORKBOOK_ARCHIVE_INVALID"}


def test_wrong_headers_fail_closed():
    content = make_change(sample_workbook(), lambda book: setattr(book["Accounts"]["A1"], "value", "other"))
    with pytest.raises(CoverageError) as exc:
        preview(content)
    assert exc.value.code == "WORKBOOK_INVALID"


@pytest.mark.parametrize("alter", [
    lambda book: setattr(book["Mapping"].row_dimensions[2], "hidden", True),
    lambda book: setattr(book["Accounts"].column_dimensions["Q"], "hidden", True),
    lambda book: book["Mapping"].merge_cells("C2:D2"),
    lambda book: setattr(book["Notes"], "sheet_state", "hidden"),
])
def test_hidden_or_merged_inputs_fail_closed(alter):
    content = make_change(sample_workbook(), alter)
    with pytest.raises(CoverageError) as exc:
        preview(content)
    # The existing Pilot preflight may reject a merge before the stronger
    # non-streaming visibility check. Both paths fail closed.
    assert exc.value.code in {"HIDDEN_INPUT", "MERGED_INPUT", "WORKBOOK_INVALID"}


def test_original_bytes_never_mutated():
    source = sample_workbook()
    before = bytes(source)
    result = preview(source)
    assert source == before
    assert len(result.source_sha256) == 64


def test_result_serializable_without_float_amounts():
    import json
    result = preview(sample_workbook())
    encoded = json.dumps(result.to_dict(), sort_keys=True)
    assert '"110.25"' in encoded
    assert '"releasable": false' in encoded


def test_rejects_nonbytes_and_large_inputs():
    with pytest.raises(CoverageError):
        preview_mapping_coverage("not-bytes", company_code="001", fiscal_year=1404)
    with pytest.raises(CoverageError):
        preview(b"x" * (20 * 1024 * 1024 + 1))


def test_empty_scoped_accounts_fail_closed():
    content = make_change(sample_workbook(), lambda book: (
        book["Accounts"].delete_rows(2, 2),
        book["Mapping"].delete_rows(2, 1),
    ))
    with pytest.raises(CoverageError) as exc:
        preview(content)
    assert exc.value.code == "NO_ACCOUNTS_IN_SCOPE"
