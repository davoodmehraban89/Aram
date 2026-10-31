"""Safety and content tests for the nonofficial local analyst workbook."""
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from aram_backend.analyst_review_workbook import (
    DISCLAIMER, build_review_workbook_bytes, write_review_workbook,
)
from aram_backend.coverage_preview import CoverageError
from aram_backend.synthetic_five_statement_demo import cash_transactions
from test_five_statement_pipeline_integration import build_pipeline


def test_review_pack_contains_exact_strings_all_statements_and_disclaimers(tmp_path):
    result = build_pipeline()
    path = write_review_workbook(tmp_path / "review-only.xlsx", result)
    assert path.is_file()
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == [
            "اطلاعات", "خلاصه مبالغ", "ترازنامه", "سود و زیان",
            "حقوق مالکانه", "جریان نقد", "تراز آزمایشی", "کنترل‌ها",
        ]
        for sheet in workbook.worksheets:
            assert DISCLAIMER in sheet["A1"].value
        summary = list(workbook["خلاصه مبالغ"].iter_rows(values_only=True))
        assert ("سود و زیان", "سود خالص", "20", "پیش‌نویس قابل بررسی؛ غیررسمی") in summary
        balance = list(workbook["ترازنامه"].iter_rows(values_only=True))
        assert any(row[1] == "Cash and cash equivalents" and row[8] == "50" for row in balance)
        trial = list(workbook["تراز آزمایشی"].iter_rows(values_only=True))
        assert len(trial) == len(result.trial_balance.accounts) + 4
        assert all(
            cell.data_type != "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()


def test_review_pack_never_overwrites_existing_files(tmp_path):
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"do not overwrite")
    with pytest.raises(CoverageError, match="TARGET_EXISTS"):
        write_review_workbook(target, build_pipeline())
    assert target.read_bytes() == b"do not overwrite"


def test_review_pack_refuses_wrong_extension_and_missing_parent(tmp_path):
    result = build_pipeline()
    with pytest.raises(CoverageError, match="EXTENSION"):
        write_review_workbook(tmp_path / "review.xls", result)
    with pytest.raises(CoverageError, match="PARENT"):
        write_review_workbook(tmp_path / "missing" / "review.xlsx", result)


def test_formula_like_user_text_is_written_as_a_literal_string(tmp_path):
    result = build_pipeline()
    line = replace(
        result.balance_sheet.lines[0],
        line_label='=HYPERLINK("https://example.invalid","open")',
    )
    changed = replace(
        result, balance_sheet=replace(
            result.balance_sheet, lines=(line,) + result.balance_sheet.lines[1:],
        ),
    )
    path = write_review_workbook(tmp_path / "literal.xlsx", changed)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = list(workbook["ترازنامه"].iter_rows())
        matching = [
            cell for row in rows for cell in row
            if cell.value == '=HYPERLINK("https://example.invalid","open")'
        ]
        assert len(matching) == 1
        assert matching[0].data_type == "s"
        assert all(
            cell.data_type != "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()


def test_blocked_pipeline_pack_is_explicitly_diagnostic(tmp_path):
    result = build_pipeline(transactions=cash_transactions()[:-1])
    assert not result.eligible_for_analyst_review
    content = build_review_workbook_bytes(result)
    assert content.startswith(b"PK")
    path = write_review_workbook(tmp_path / "blocked.xlsx", result)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        summary = list(workbook["خلاصه مبالغ"].iter_rows(values_only=True))
        assert any("مسدود" in str(row[3]) for row in summary[4:])
        info = list(workbook["اطلاعات"].iter_rows(values_only=True))
        assert any(row[0] == "وضعیت کلی" and "مسدود" in row[1] for row in info)
    finally:
        workbook.close()


def test_review_pack_rejects_bidi_and_control_text():
    result = build_pipeline()
    line = replace(result.income_statement.lines[0], line_label="Revenue\u202eforged")
    changed = replace(
        result, income_statement=replace(
            result.income_statement, lines=(line,) + result.income_statement.lines[1:],
        ),
    )
    with pytest.raises(CoverageError):
        build_review_workbook_bytes(changed)


def test_review_pack_refuses_symlink_targets(tmp_path):
    target = tmp_path / "real.xlsx"
    link = tmp_path / "alias.xlsx"
    try:
        link.symlink_to(target)
    except (OSError, NotImplementedError):
        pytest.skip("Symlinks unavailable")
    with pytest.raises(CoverageError, match="TARGET_EXISTS"):
        write_review_workbook(link, build_pipeline())
