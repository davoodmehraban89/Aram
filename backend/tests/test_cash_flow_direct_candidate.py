"""Synthetic-only direct cash flow candidate tests; no customer data."""
from __future__ import annotations

from dataclasses import replace
from io import BytesIO
import json

import pytest
from openpyxl import Workbook, load_workbook

from aram_backend.cash_flow_direct_candidate import (
    CashAccountPolicy, CashCounterpartPolicy, CashFlowEntry,
    build_direct_cash_flow_candidate,
)
from aram_backend.coverage_preview import CoverageError
from aram_backend.pilot_workbook import SHEET_HEADERS


ACCOUNT_ROWS = (
    ("CASH1", "Cash account one", "100", "40", "40", "100"),
    ("CASH2", "Cash account two", "20", "20", "10", "30"),
    ("REV", "Sales counterpart", "0", "0", "40", "-40"),
    ("EXP1", "Operating expense", "0", "20", "0", "20"),
    ("ASSET", "Equipment", "0", "10", "0", "10"),
    ("CAP", "Capital", "0", "0", "10", "-10"),
    ("EXP2", "Other operating expense", "0", "10", "0", "10"),
)


def workbook(*, mutate=None):
    book = Workbook()
    book.remove(book.active)
    for name, headers in SHEET_HEADERS.items():
        book.create_sheet(name).append(list(headers))

    def add(name, **values):
        book[name].append([values.get(key) for key in SHEET_HEADERS[name]])

    add("Company", company_code="001", company_name="Synthetic only",
        fiscal_year=1404, currency="IRR", unit="Rial", status="Draft")
    for index, (code, name, opening, debit, credit, closing) in enumerate(ACCOUNT_ROWS, 2):
        add("Accounts", company_code="001", fiscal_year=1404,
            general_code=code, general_name=name, account_full_code=code,
            account_name=name, opening=opening, debit=debit, credit=credit,
            closing=closing, source_row=index + 20, row_status="Draft",
            closing_delta="0")
    if mutate:
        mutate(book)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def cash_accounts():
    return [
        CashAccountPolicy("CASH1", "SYNTH-CASH-POLICY-1"),
        CashAccountPolicy("CASH2", "SYNTH-CASH-POLICY-2"),
    ]


def counterpart_policies():
    return [
        CashCounterpartPolicy("REV", "OPERATING", "SYNTH-REV"),
        CashCounterpartPolicy("EXP1", "OPERATING", "SYNTH-EXP1"),
        CashCounterpartPolicy("ASSET", "INVESTING", "SYNTH-ASSET"),
        CashCounterpartPolicy("CAP", "FINANCING", "SYNTH-CAP"),
        CashCounterpartPolicy("EXP2", "OPERATING", "SYNTH-EXP2"),
    ]


def transactions():
    common = {"company_code": "001", "fiscal_year": 1404}
    return [
        CashFlowEntry(**common, line_id="L1", transaction_id="V1",
            posting_date="1404-01-01", cash_account_code="CASH1",
            counterpart_account_code="REV", cash_debit="40", cash_credit="0",
            counterpart_debit="0", counterpart_credit="40", evidence_ref="SYNTH-T1"),
        CashFlowEntry(**common, line_id="L2", transaction_id="V2",
            posting_date="1404-01-02", cash_account_code="CASH1",
            counterpart_account_code="EXP1", cash_debit="0", cash_credit="20",
            counterpart_debit="20", counterpart_credit="0", evidence_ref="SYNTH-T2"),
        CashFlowEntry(**common, line_id="L3", transaction_id="V3",
            posting_date="1404-01-03", cash_account_code="CASH1",
            counterpart_account_code="ASSET", cash_debit="0", cash_credit="10",
            counterpart_debit="10", counterpart_credit="0", evidence_ref="SYNTH-T3"),
        CashFlowEntry(**common, line_id="L4", transaction_id="V4",
            posting_date="1404-01-04", cash_account_code="CASH1",
            counterpart_account_code="CASH2", cash_debit="0", cash_credit="10",
            counterpart_debit="10", counterpart_credit="0", evidence_ref="SYNTH-T4"),
        CashFlowEntry(**common, line_id="L5", transaction_id="V5",
            posting_date="1404-01-05", cash_account_code="CASH2",
            counterpart_account_code="CAP", cash_debit="10", cash_credit="0",
            counterpart_debit="0", counterpart_credit="10", evidence_ref="SYNTH-T5"),
        CashFlowEntry(**common, line_id="L6", transaction_id="V6",
            posting_date="1404-01-06", cash_account_code="CASH2",
            counterpart_account_code="EXP2", cash_debit="0", cash_credit="10",
            counterpart_debit="10", counterpart_credit="0", evidence_ref="SYNTH-T6"),
    ]


def build(*, data=None, cash=None, policies=None, entries=None,
          company="001", year=1404, currency="IRR", unit="Rial"):
    return build_direct_cash_flow_candidate(
        workbook() if data is None else data,
        company_code=company, fiscal_year=year, currency=currency, unit=unit,
        cash_accounts=cash_accounts() if cash is None else cash,
        counterpart_policies=counterpart_policies() if policies is None else policies,
        transactions=transactions() if entries is None else entries,
    )


def edit(data, change):
    book = load_workbook(BytesIO(data))
    change(book)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_direct_cash_flow_classification_and_cash_rollforward():
    result = build()
    assert result.eligible_for_analyst_review
    assert not result.findings
    assert result.opening_cash_and_equivalents == "120"
    assert result.operating_inflows == "40"
    assert result.operating_outflows == "30"
    assert result.operating_net_cash_flow == "10"
    assert result.investing_inflows == "0"
    assert result.investing_outflows == "10"
    assert result.investing_net_cash_flow == "-10"
    assert result.financing_inflows == "10"
    assert result.financing_outflows == "0"
    assert result.financing_net_cash_flow == "10"
    assert result.internal_transfer_inflows == "10"
    assert result.internal_transfer_outflows == "10"
    assert result.net_change_in_cash == "10"
    assert result.closing_cash_and_equivalents == "130"
    assert result.rollforward_difference == "0"
    assert result.cash_account_count == 2
    assert result.transaction_count == 6
    assert len(result.cash_flow_lines) == 6
    assert all(a.debit_difference == "0" and a.credit_difference == "0"
               for a in result.account_reconciliations)
    assert result.kind == "UNAPPROVED_DIRECT_CASH_FLOW_CANDIDATE"
    assert not result.cash_account_scope_owner_approved
    assert not result.transaction_ledger_completeness_verified
    assert not result.accounting_policy_authorized
    assert not result.official_financial_statement
    assert not result.independently_reviewed
    assert not result.export_authorized
    assert not result.releasable


def test_missing_transaction_line_blocks_turnover_reconciliation():
    result = build(entries=transactions()[:-1])
    assert not result.eligible_for_analyst_review
    assert "CASH_ACCOUNT_CREDIT_TURNOVER_MISMATCH" in result.findings


def test_missing_counterpart_classification_blocks_candidate():
    policies = counterpart_policies()[:-1]
    result = build(policies=policies)
    assert not result.eligible_for_analyst_review
    assert "MISSING_COUNTERPART_ACTIVITY_POLICY" in result.findings
    assert "CASH_FLOW_CASH_BALANCE_ROLLFORWARD_MISMATCH" in result.findings


def test_orphan_counterpart_policy_blocks_candidate():
    policies = counterpart_policies() + [
        CashCounterpartPolicy("UNUSED", "OPERATING", "SYNTH-UNUSED")
    ]
    result = build(policies=policies)
    assert not result.eligible_for_analyst_review
    assert "ORPHAN_COUNTERPART_ACTIVITY_POLICY" in result.findings


def test_internal_cash_transfer_is_excluded_from_activity_categories():
    result = build()
    assert result.internal_transfer_inflows == result.internal_transfer_outflows
    assert result.operating_net_cash_flow == "10"
    assert result.investing_net_cash_flow == "-10"
    assert result.financing_net_cash_flow == "10"


def test_unbalanced_cash_counterpart_pair_rejected():
    entries = transactions()
    entries[0] = replace(entries[0], counterpart_credit="39")
    with pytest.raises(CoverageError) as error:
        build(entries=entries)
    assert error.value.code == "CASH_FLOW_COUNTERPART_NOT_BALANCED"


def test_duplicate_line_id_rejected():
    entries = transactions()
    entries[1] = replace(entries[1], line_id="L1")
    with pytest.raises(CoverageError) as error:
        build(entries=entries)
    assert error.value.code == "DUPLICATE_CASH_FLOW_LINE_ID"


def test_wrong_company_or_year_transaction_rejected():
    entries = transactions()
    entries[0] = replace(entries[0], company_code="002")
    with pytest.raises(CoverageError) as error:
        build(entries=entries)
    assert error.value.code == "CASH_FLOW_TRANSACTION_SCOPE_MISMATCH"
    entries = transactions()
    entries[0] = replace(entries[0], fiscal_year=1403)
    with pytest.raises(CoverageError):
        build(entries=entries)


def test_invalid_posting_date_rejected():
    entries = transactions()
    entries[0] = replace(entries[0], posting_date="1404-13-01")
    with pytest.raises(CoverageError) as error:
        build(entries=entries)
    assert error.value.code == "INVALID_CASH_FLOW_POSTING_DATE"


def test_currency_or_unit_mismatch_rejected():
    with pytest.raises(CoverageError) as error:
        build(currency="USD")
    assert error.value.code == "CASH_FLOW_CURRENCY_UNIT_MISMATCH"
    with pytest.raises(CoverageError):
        build(unit="Thousand Rial")


def test_cash_account_missing_from_source_rejected():
    with pytest.raises(CoverageError) as error:
        build(cash=[CashAccountPolicy("MISSING", "SYNTH")])
    assert error.value.code == "CASH_ACCOUNT_NOT_FOUND_IN_SOURCE"


def test_transaction_unknown_cash_or_counterpart_account_rejected():
    entries = transactions()
    entries[0] = replace(entries[0], cash_account_code="UNKNOWN")
    with pytest.raises(CoverageError):
        build(entries=entries)
    entries = transactions()
    entries[0] = replace(entries[0], counterpart_account_code="UNKNOWN")
    with pytest.raises(CoverageError):
        build(entries=entries)


def test_negative_or_two_sided_cash_posting_rejected():
    entries = transactions()
    entries[0] = replace(entries[0], cash_debit="-40")
    with pytest.raises(CoverageError):
        build(entries=entries)
    entries = transactions()
    entries[0] = replace(entries[0], cash_credit="1")
    with pytest.raises(CoverageError):
        build(entries=entries)


def test_duplicate_cash_or_counterpart_policies_rejected():
    with pytest.raises(CoverageError):
        build(cash=cash_accounts() + [cash_accounts()[0]])
    with pytest.raises(CoverageError):
        build(policies=counterpart_policies() + [counterpart_policies()[0]])


def test_invalid_activity_or_cash_counterpart_policy_rejected():
    bad = counterpart_policies()
    bad[0] = replace(bad[0], activity="CASH_TRANSFER")
    with pytest.raises(CoverageError):
        build(policies=bad)
    bad = counterpart_policies() + [
        CashCounterpartPolicy("CASH1", "OPERATING", "SYNTH-BAD")
    ]
    with pytest.raises(CoverageError):
        build(policies=bad)


def test_all_zero_cash_balances_and_no_transactions_are_valid():
    def zero(book):
        for row in book["Accounts"].iter_rows(min_row=2):
            for col in (16, 17, 18, 19):
                row[col].value = "0"
    data = edit(workbook(), zero)
    result = build(data=data, entries=[])
    assert result.eligible_for_analyst_review
    assert result.opening_cash_and_equivalents == "0"
    assert result.net_change_in_cash == "0"
    assert result.closing_cash_and_equivalents == "0"


def test_source_cash_rollforward_invalid_rejected():
    data = edit(workbook(), lambda book: setattr(book["Accounts"]["T2"], "value", "101"))
    with pytest.raises(CoverageError):
        build(data=data)


def test_wrong_company_year_and_currency_source_scope_rejected():
    with pytest.raises(CoverageError):
        build(company="999")
    with pytest.raises(CoverageError):
        build(year=1403)
    data = edit(workbook(), lambda book: setattr(book["Company"]["D2"], "value", "USD"))
    with pytest.raises(CoverageError):
        build(data=data)


def test_policy_and_transaction_hashes_are_stable_and_source_unchanged():
    data = workbook()
    original = bytes(data)
    first = build(data=data)
    second = build(data=data, entries=list(reversed(transactions())))
    assert data == original
    assert first.transaction_ledger_sha256 == second.transaction_ledger_sha256
    assert first.classification_policies_sha256 == second.classification_policies_sha256
    assert len(first.source_sha256) == 64
    assert not first.releasable


def test_input_size_and_transaction_count_are_bounded():
    with pytest.raises(CoverageError):
        build(data=b"x" * (20 * 1024 * 1024 + 1))
    with pytest.raises(CoverageError):
        build(entries=transactions() * 5001)
