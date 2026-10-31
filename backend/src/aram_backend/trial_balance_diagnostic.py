"""Read-only trial balance diagnostic for an authorized Pilot XLSX snapshot.

This checks source account arithmetic and debit/credit equality within an
explicit company-year scope. It does not perform financial postings, issue
audited statements, or assert that the provided ledger is complete.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from hashlib import sha256
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from aram_backend.coverage_preview import (
    CoverageError, MAX_INPUT_BYTES, _assert_visible, _guard_archive,
)
from aram_backend.pilot_workbook import SHEET_HEADERS, preflight_pilot_workbook
from aram_backend.reconciliation_review import _money, _sum_exact


@dataclass(frozen=True)
class TrialBalanceAccount:
    account_code: str
    account_name: str
    excel_row: int
    source_row: int
    opening: str
    debit: str
    credit: str
    closing: str
    debit_net_movement: str


@dataclass(frozen=True)
class TrialBalanceDiagnostic:
    company_code: str
    fiscal_year: int
    source_sha256: str
    account_count: int
    opening_debit_net_total: str
    debit_turnover_total: str
    credit_turnover_total: str
    closing_debit_net_total: str
    movement_difference: str
    opening_to_closing_difference: str
    findings: tuple[str, ...]
    accounts: tuple[TrialBalanceAccount, ...]
    balanced_for_analyst_review: bool
    kind: str = "UNAPPROVED_TRIAL_BALANCE_DIAGNOSTIC"
    source_ledger_completeness_verified: bool = False
    accounting_policies_verified: bool = False
    official_financial_statement: bool = False
    independently_reviewed: bool = False
    releasable: bool = False

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _scoped(row: dict[str, Any], *, company: str, fiscal_year: int) -> bool:
    return (
        isinstance(row.get("company_code"), str)
        and row["company_code"].strip() == company
        and type(row.get("fiscal_year")) is int
        and row["fiscal_year"] == fiscal_year
    )


def _read_records(sheet: Any, columns: tuple[str, ...]) -> list[tuple[int, dict[str, Any]]]:
    records: list[tuple[int, dict[str, Any]]] = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, values_only=True), 2
    ):
        row = dict(zip(columns, values))
        if any(cell is not None and str(cell).strip() for cell in row.values()):
            records.append((row_number, row))
    return records


def analyze_trial_balance(
    workbook_bytes: bytes, *, company_code: str, fiscal_year: int,
) -> TrialBalanceDiagnostic:
    """Reconcile row and column totals without claiming source completeness.

    The check assumes opening and closing figures are *debit-net*, meaning a
    credit-net account has a negative balance. It does not infer the sign from
    an account label. Incorrectly signed input can appear balanced, so the
    result is NEVER sufficient for financial authorization or release.
    """
    if (not isinstance(company_code, str) or not company_code.strip()
            or len(company_code) > 64):
        raise CoverageError("TRIAL_BALANCE_COMPANY_SCOPE_INVALID")
    if type(fiscal_year) is not int or fiscal_year <= 0:
        raise CoverageError("TRIAL_BALANCE_YEAR_SCOPE_INVALID")
    if (not isinstance(workbook_bytes, bytes)
            or not 0 < len(workbook_bytes) <= MAX_INPUT_BYTES):
        raise CoverageError("TRIAL_BALANCE_SIZE_INVALID")
    _guard_archive(workbook_bytes)
    try:
        preflight = preflight_pilot_workbook(workbook_bytes)
    except Exception as exc:
        raise CoverageError("TRIAL_BALANCE_WORKBOOK_INVALID") from exc
    if not preflight.valid:
        raise CoverageError("TRIAL_BALANCE_WORKBOOK_INVALID")

    try:
        book = load_workbook(
            BytesIO(workbook_bytes), read_only=False,
            data_only=False, keep_links=False,
        )
    except Exception as exc:
        raise CoverageError("TRIAL_BALANCE_WORKBOOK_INVALID") from exc
    try:
        for sheet in book.worksheets:
            _assert_visible(sheet)
        for name, expected in SHEET_HEADERS.items():
            if name not in book or tuple(cell.value for cell in book[name][1]) != expected:
                raise CoverageError("TRIAL_BALANCE_WORKBOOK_INVALID")

        company = company_code.strip()
        scope = {"company": company, "fiscal_year": fiscal_year}
        companies = [
            row for _, row in _read_records(book["Company"], SHEET_HEADERS["Company"])
            if _scoped(row, **scope)
        ]
        if len(companies) != 1:
            raise CoverageError("TRIAL_BALANCE_COMPANY_NOT_FOUND")

        accounts: list[TrialBalanceAccount] = []
        openings = []
        debits = []
        credits = []
        closings = []
        net_movements = []
        codes: set[str] = set()

        for row_number, row in _read_records(book["Accounts"], SHEET_HEADERS["Accounts"]):
            if not _scoped(row, **scope):
                continue
            code = row.get("account_full_code")
            if not isinstance(code, str) or not code.strip():
                raise CoverageError("TRIAL_BALANCE_ACCOUNT_CODE_INVALID")
            code = code.strip()
            if code in codes:
                raise CoverageError("TRIAL_BALANCE_DUPLICATE_ACCOUNT")
            codes.add(code)
            opening, debit, credit, closing = (
                _money(str(row[key])) for key in ("opening", "debit", "credit", "closing")
            )
            movement = _sum_exact([debit, -credit])
            if _sum_exact([opening, movement, -closing]) != 0:
                raise CoverageError("TRIAL_BALANCE_ROW_DELTA_INVALID")
            if debit < 0 or credit < 0:
                raise CoverageError("TRIAL_BALANCE_NEGATIVE_TURNOVER")
            if row.get("source_row") is None or type(row["source_row"]) is not int:
                raise CoverageError("TRIAL_BALANCE_SOURCE_ROW_INVALID")
            accounts.append(TrialBalanceAccount(
                account_code=code, account_name=str(row["account_name"]),
                excel_row=row_number, source_row=row["source_row"],
                opening=str(opening), debit=str(debit), credit=str(credit),
                closing=str(closing), debit_net_movement=str(movement),
            ))
            openings.append(opening)
            debits.append(debit)
            credits.append(credit)
            closings.append(closing)
            net_movements.append(movement)

        if not accounts:
            raise CoverageError("TRIAL_BALANCE_NO_ACCOUNTS")
        opening_total = _sum_exact(openings)
        debit_total = _sum_exact(debits)
        credit_total = _sum_exact(credits)
        closing_total = _sum_exact(closings)
        net_total = _sum_exact(net_movements)
        movement_diff = _sum_exact([debit_total, -credit_total])
        rolling_diff = _sum_exact([opening_total, net_total, -closing_total])
        findings: list[str] = []
        if opening_total != 0:
            findings.append("OPENING_DEBIT_NET_UNBALANCED")
        if movement_diff != 0:
            findings.append("DEBIT_CREDIT_TURNOVER_UNBALANCED")
        if closing_total != 0:
            findings.append("CLOSING_DEBIT_NET_UNBALANCED")
        if rolling_diff != 0:
            findings.append("ROLL_FORWARD_INCONSISTENT")

        return TrialBalanceDiagnostic(
            company_code=company, fiscal_year=fiscal_year,
            source_sha256=sha256(workbook_bytes).hexdigest(),
            account_count=len(accounts),
            opening_debit_net_total=str(opening_total),
            debit_turnover_total=str(debit_total),
            credit_turnover_total=str(credit_total),
            closing_debit_net_total=str(closing_total),
            movement_difference=str(movement_diff),
            opening_to_closing_difference=str(rolling_diff),
            findings=tuple(findings),
            accounts=tuple(sorted(accounts, key=lambda acc: acc.account_code)),
            balanced_for_analyst_review=not findings,
        )
    finally:
        book.close()
