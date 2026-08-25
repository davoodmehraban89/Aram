from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


SHEET_HEADERS: dict[str, tuple[str, ...]] = {
    "Company": (
        "company_code", "company_name", "fiscal_year", "currency", "unit",
        "prepared_by", "status",
    ),
    "Accounts": (
        "company_code", "fiscal_year", "general_code", "general_name",
        "subsidiary_code", "subsidiary_name", "detail1_code", "detail1_name",
        "detail2_code", "detail2_name", "detail3_code", "detail3_name",
        "detail4_code", "detail4_name", "account_full_code", "account_name",
        "opening", "debit", "credit", "closing", "source_row", "row_status",
        "closing_delta",
    ),
    "Notes": (
        "company_code", "fiscal_year", "statement", "statement_section",
        "note_number", "note_title", "subnote_code", "subnote_title",
    ),
    "Mapping": (
        "mapping_id", "company_code", "fiscal_year", "account_full_code",
        "account_level", "general_code", "subsidiary_code", "detail1_code",
        "detail2_code", "detail3_code", "detail4_code", "target_statement",
        "statement_section", "note_number", "subnote_code", "mapping_version",
        "approval_status", "reviewer", "review_date", "source", "notes",
        "validation_status",
    ),
    "Test_Log": (
        "tester", "feature", "expected_result", "actual_result", "status",
        "severity", "evidence", "notes",
    ),
    "Checks": ("check_code", "check_name", "status", "affected_rows", "details"),
    "Guide_FA": ("section", "instruction"),
    "Lists": ("list_name", "value"),
}

_SHEET_ORDER = {name: position for position, name in enumerate(SHEET_HEADERS)}
_CODE_FIELDS = {
    "company_code", "general_code", "subsidiary_code", "detail1_code",
    "detail2_code", "detail3_code", "detail4_code", "account_full_code",
    "note_number", "subnote_code", "mapping_id",
}
_REQUIRED: dict[str, set[str]] = {
    "Company": {"company_code", "company_name", "fiscal_year", "status"},
    "Accounts": {
        "company_code", "fiscal_year", "general_code", "general_name",
        "account_full_code", "account_name", "opening", "debit", "credit",
        "closing", "source_row", "row_status",
    },
    "Notes": {"company_code", "fiscal_year", "statement", "note_number", "note_title"},
    "Mapping": {
        "mapping_id", "company_code", "fiscal_year", "account_full_code",
        "account_level", "mapping_version", "approval_status",
    },
}


@dataclass(frozen=True)
class PilotIssue:
    severity: str
    code: str
    sheet: str
    row: int
    field: str
    message: str


@dataclass(frozen=True)
class PilotPreflightResult:
    valid: bool
    issues: tuple[PilotIssue, ...]
    issue_counts: dict[str, int]


def _empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _issue(issues: list[PilotIssue], code: str, sheet: str, row: int,
           field: str, message: str, severity: str = "ERROR") -> None:
    issues.append(PilotIssue(severity, code, sheet, row, field, message))


def _text_code(value: Any, issues: list[PilotIssue], sheet: str, row: int,
               field: str) -> str | None:
    if _empty(value):
        return None
    if not isinstance(value, str):
        _issue(issues, "CODE_NOT_TEXT", sheet, row, field,
               "Code cells must use Excel text format; numeric coercion is forbidden.")
        return None
    return value.strip()


def _positive_int(value: Any, issues: list[PilotIssue], sheet: str, row: int,
                  field: str) -> int | None:
    if _empty(value):
        return None
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        _issue(issues, "INVALID_POSITIVE_INTEGER", sheet, row, field,
               "Value must be a positive integer.")
        return None
    return value


def _finite_decimal(value: Any, issues: list[PilotIssue], sheet: str, row: int,
                    field: str) -> Decimal | None:
    if _empty(value):
        return None
    if isinstance(value, bool) or isinstance(value, str) and value.startswith("="):
        _issue(issues, "INVALID_FINANCIAL_VALUE", sheet, row, field,
               "Financial values must be finite numbers, not booleans or formulas.")
        return None
    try:
        parsed = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        _issue(issues, "INVALID_FINANCIAL_VALUE", sheet, row, field,
               "Financial value is not a valid decimal number.")
        return None
    if not parsed.is_finite():
        _issue(issues, "NON_FINITE_FINANCIAL_VALUE", sheet, row, field,
               "NaN and Infinity are not permitted.")
        return None
    return parsed


def _rows(sheet: Any, headers: tuple[str, ...]) -> list[tuple[int, dict[str, Any]]]:
    result: list[tuple[int, dict[str, Any]]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        row = dict(zip(headers, values))
        if any(not _empty(value) for value in row.values()):
            result.append((row_number, row))
    return result


def preflight_pilot_workbook(workbook_bytes: bytes) -> PilotPreflightResult:
    """Validate a Pilot workbook fully in memory without persistence or mutation."""
    issues: list[PilotIssue] = []
    try:
        workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        issue = PilotIssue("ERROR", "INVALID_WORKBOOK", "__workbook__", 0, "__file__", str(exc))
        return PilotPreflightResult(False, (issue,), {"ERROR": 1})

    valid_sheets: set[str] = set()
    for sheet_name, expected in SHEET_HEADERS.items():
        if sheet_name not in workbook.sheetnames:
            _issue(issues, "MISSING_SHEET", sheet_name, 1, "__sheet__",
                   f"Required sheet {sheet_name!r} is missing.")
            continue
        actual = tuple(cell.value for cell in workbook[sheet_name][1])
        if actual != expected:
            _issue(issues, "HEADER_MISMATCH", sheet_name, 1, "__headers__",
                   f"Expected exact header order: {', '.join(expected)}.")
            continue
        valid_sheets.add(sheet_name)

    parsed: dict[str, list[tuple[int, dict[str, Any]]]] = {
        name: _rows(workbook[name], SHEET_HEADERS[name])
        for name in valid_sheets
    }
    workbook.close()

    for sheet_name in ("Company", "Accounts", "Notes", "Mapping"):
        for row_number, row in parsed.get(sheet_name, []):
            for field in _REQUIRED[sheet_name]:
                if _empty(row[field]):
                    _issue(issues, "REQUIRED_FIELD", sheet_name, row_number, field,
                           "Required field is empty.")
            for field in _CODE_FIELDS.intersection(row):
                _text_code(row[field], issues, sheet_name, row_number, field)
            _positive_int(row.get("fiscal_year"), issues, sheet_name, row_number, "fiscal_year")

    companies: set[tuple[str, int]] = set()
    for row_number, row in parsed.get("Company", []):
        code = _text_code(row["company_code"], [], "Company", row_number, "company_code")
        year = row["fiscal_year"] if isinstance(row["fiscal_year"], int) and not isinstance(row["fiscal_year"], bool) else None
        if code and year and year > 0:
            key = (code, year)
            if key in companies:
                _issue(issues, "DUPLICATE_COMPANY_YEAR", "Company", row_number,
                       "company_code", "Company/fiscal-year key must be unique.")
            companies.add(key)
        if not _empty(row["status"]) and row["status"] not in {"Draft", "Approved", "Locked"}:
            _issue(issues, "INVALID_LIFECYCLE", "Company", row_number, "status",
                   "Allowed values are Draft, Approved, and Locked.")

    account_keys: set[tuple[str, int, str]] = set()
    for row_number, row in parsed.get("Accounts", []):
        code = row["company_code"] if isinstance(row["company_code"], str) else None
        year = row["fiscal_year"] if isinstance(row["fiscal_year"], int) and not isinstance(row["fiscal_year"], bool) else None
        account = row["account_full_code"] if isinstance(row["account_full_code"], str) else None
        if code and year and (code.strip(), year) not in companies:
            _issue(issues, "UNKNOWN_COMPANY_YEAR", "Accounts", row_number, "company_code",
                   "Referenced company/fiscal year is not defined in Company.")
        if code and year and account:
            key = (code.strip(), year, account.strip())
            if key in account_keys:
                _issue(issues, "DUPLICATE_ACCOUNT", "Accounts", row_number,
                       "account_full_code", "Account key is duplicated within company/fiscal year.")
            account_keys.add(key)
        values = [_finite_decimal(row[field], issues, "Accounts", row_number, field)
                  for field in ("opening", "debit", "credit", "closing")]
        _positive_int(row["source_row"], issues, "Accounts", row_number, "source_row")
        if not _empty(row["row_status"]) and row["row_status"] not in {"Draft", "Reviewed", "Approved"}:
            _issue(issues, "INVALID_LIFECYCLE", "Accounts", row_number, "row_status",
                   "Allowed values are Draft, Reviewed, and Approved.")
        if all(value is not None for value in values):
            delta = values[0] + values[1] - values[2] - values[3]
            if delta != 0:
                _issue(issues, "NON_ZERO_CLOSING_DELTA", "Accounts", row_number,
                       "closing_delta", f"opening + debit - credit - closing equals {delta}.")

    for sheet_name in ("Notes", "Mapping"):
        for row_number, row in parsed.get(sheet_name, []):
            code = row["company_code"] if isinstance(row["company_code"], str) else None
            year = row["fiscal_year"] if isinstance(row["fiscal_year"], int) and not isinstance(row["fiscal_year"], bool) else None
            if code and year and (code.strip(), year) not in companies:
                _issue(issues, "UNKNOWN_COMPANY_YEAR", sheet_name, row_number,
                       "company_code", "Referenced company/fiscal year is not defined in Company.")

    mapping_ids: set[str] = set()
    for row_number, row in parsed.get("Mapping", []):
        mapping_id = row["mapping_id"] if isinstance(row["mapping_id"], str) else None
        if mapping_id:
            if mapping_id.strip() in mapping_ids:
                _issue(issues, "DUPLICATE_MAPPING_ID", "Mapping", row_number,
                       "mapping_id", "Mapping ID must be unique.")
            mapping_ids.add(mapping_id.strip())
        _positive_int(row["mapping_version"], issues, "Mapping", row_number, "mapping_version")
        if not _empty(row["approval_status"]) and row["approval_status"] not in {"Draft", "Reviewed", "Approved"}:
            _issue(issues, "INVALID_LIFECYCLE", "Mapping", row_number, "approval_status",
                   "Allowed values are Draft, Reviewed, and Approved.")
        code = row["company_code"] if isinstance(row["company_code"], str) else None
        year = row["fiscal_year"] if isinstance(row["fiscal_year"], int) and not isinstance(row["fiscal_year"], bool) else None
        account = row["account_full_code"] if isinstance(row["account_full_code"], str) else None
        if code and year and account and (code.strip(), year, account.strip()) not in account_keys:
            _issue(issues, "UNKNOWN_MAPPING_ACCOUNT", "Mapping", row_number,
                   "account_full_code", "Mapped account does not exist in Accounts for this company/fiscal year.")

    ordered = tuple(sorted(issues, key=lambda item: (
        _SHEET_ORDER.get(item.sheet, -1), item.row,
        SHEET_HEADERS.get(item.sheet, ()).index(item.field)
        if item.field in SHEET_HEADERS.get(item.sheet, ()) else -1,
        item.code, item.message,
    )))
    counts = dict(sorted(Counter(item.severity for item in ordered).items()))
    return PilotPreflightResult(not any(item.severity == "ERROR" for item in ordered), ordered, counts)
