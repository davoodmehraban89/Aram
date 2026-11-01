from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from aram_backend.workbook_inspection import inspect_xlsx


MAX_MAPPING_ROWS = 100_000

_REQUIRED_HEADERS = (
    "کد حساب",
    "نوع",
    "کد یادداشت",
    "شرح یادداشت",
    "فرمول یادداشت اصلی",
    "فرمول یادداشت فرعی",
)


@dataclass(frozen=True)
class MappingSourceRow:
    """One untouched source mapping row with its original Excel row number."""

    row_number: int
    account_code: str
    mapping_type: str
    note_code: str
    note_description: str
    main_formula: str
    sub_formula: str
    source_values: tuple[tuple[str, Any], ...]


class MappingSourceParseError(ValueError):
    """Raised when a mapping workbook row cannot be read unambiguously."""


def parse_mapping_source(content: bytes) -> Iterator[MappingSourceRow]:
    """Stream mapping rows without evaluating formulas or inventing precedence.

    All source columns are retained in source_values so later matching logic can
    explicitly interpret dimensions only after their real headers are reviewed.
    Formula cells are read as formulas (data_only=False), never executed.
    """

    inspection = inspect_xlsx(content, document_type="mapping")
    workbook = load_workbook(
        BytesIO(content), read_only=True, data_only=False, keep_links=False
    )
    try:
        worksheet = workbook[inspection.data_sheet]
        header_values = next(
            worksheet.iter_rows(
                min_row=inspection.header_row,
                max_row=inspection.header_row,
                values_only=True,
            )
        )
        headers = tuple(
            str(value).strip() if value is not None else ""
            for value in header_values
        )
        indexes = {header: index for index, header in enumerate(headers) if header}
        missing = set(_REQUIRED_HEADERS).difference(indexes)
        if missing:
            raise MappingSourceParseError(
                f"required mapping headers are missing: {sorted(missing)}"
            )

        parsed_rows = 0
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=inspection.header_row + 1,
                max_col=len(headers),
                values_only=True,
            ),
            start=inspection.header_row + 1,
        ):
            if not any(value is not None for value in values):
                continue
            parsed_rows += 1
            if parsed_rows > MAX_MAPPING_ROWS:
                raise MappingSourceParseError("mapping workbook exceeds row limit")

            def cell(header: str) -> Any:
                index = indexes[header]
                return values[index] if index < len(values) else None

            account_code = _required_code(cell("کد حساب"), row_number, "کد حساب")
            mapping_type = _required_text(cell("نوع"), row_number, "نوع")
            note_code = _required_code(cell("کد یادداشت"), row_number, "کد یادداشت")
            note_description = _optional_text(
                cell("شرح یادداشت"), row_number, "شرح یادداشت"
            )
            main_formula = _optional_text(
                cell("فرمول یادداشت اصلی"), row_number, "فرمول یادداشت اصلی"
            )
            sub_formula = _optional_text(
                cell("فرمول یادداشت فرعی"), row_number, "فرمول یادداشت فرعی"
            )
            source_values = tuple(
                (header, values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
                if header
            )
            yield MappingSourceRow(
                row_number=row_number,
                account_code=account_code,
                mapping_type=mapping_type,
                note_code=note_code,
                note_description=note_description,
                main_formula=main_formula,
                sub_formula=sub_formula,
                source_values=source_values,
            )
    finally:
        workbook.close()


def _required_text(value: object, row_number: int, field: str) -> str:
    result = _optional_text(value, row_number, field)
    if not result:
        raise MappingSourceParseError(f"row {row_number}: {field} is required")
    return result


def _optional_text(value: object, row_number: int, field: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        raise MappingSourceParseError(f"row {row_number}: {field} is not valid text")
    result = str(value).strip()
    if len(result) > 2_000:
        raise MappingSourceParseError(f"row {row_number}: {field} exceeds length limit")
    return result


def _required_code(value: object, row_number: int, field: str) -> str:
    result = _optional_text(value, row_number, field)
    if not result:
        raise MappingSourceParseError(f"row {row_number}: {field} is required")
    if isinstance(value, float) and not value.is_integer():
        raise MappingSourceParseError(f"row {row_number}: {field} must be an integer code")
    if isinstance(value, float):
        return str(int(value))
    return result
