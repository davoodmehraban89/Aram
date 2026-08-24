from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from aram_backend.workbook_inspection import inspect_xlsx


@dataclass(frozen=True)
class LedgerRecord:
    row_number: int
    account_code: str
    account_name: str
    source_year: int
    entry_type: str
    t1_code: str | None
    t2_code: str | None
    t3_code: str | None
    t4_code: str | None
    opening: Decimal
    debit: Decimal
    credit: Decimal
    closing: Decimal


class LedgerParseError(ValueError):
    """Raised when a source ledger row cannot be represented safely."""


_FIELDS = {
    "account_code": "حساب",
    "account_name": "ACCOUNT_NAME",
    "source_year": "year",
    "entry_type": "نوع",
    "t1_code": "T1_CODE",
    "t2_code": "T2_CODE",
    "t3_code": "T3_CODE",
    "t4_code": "T4_CODE",
    "opening": "OPENING",
    "debit": "Debit",
    "credit": "Credit",
    "closing": "CLOSING",
}


def parse_general_ledger(content: bytes) -> Iterator[LedgerRecord]:
    """Stream typed ledger rows without persisting or calculating values."""

    inspection = inspect_xlsx(content, document_type="general_ledger")
    workbook = load_workbook(
        BytesIO(content), read_only=True, data_only=True, keep_links=False
    )
    try:
        worksheet = workbook[inspection.data_sheet]
        header_values = next(
            worksheet.iter_rows(
                min_row=inspection.header_row,
                max_row=inspection.header_row,
                values_only=True,
            )
        )
        indexes = {
            str(value).strip(): index
            for index, value in enumerate(header_values)
            if value is not None
        }

        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=inspection.header_row + 1, values_only=True
            ),
            start=inspection.header_row + 1,
        ):
            if not any(value is not None for value in values):
                continue

            def value(field: str):
                return values[indexes[_FIELDS[field]]]

            yield LedgerRecord(
                row_number=row_number,
                account_code=_required_code(value("account_code"), row_number, "حساب"),
                account_name=_required_text(
                    value("account_name"), row_number, "ACCOUNT_NAME"
                ),
                source_year=_required_year(value("source_year"), row_number),
                entry_type=_required_text(value("entry_type"), row_number, "نوع"),
                t1_code=_optional_code(value("t1_code"), row_number, "T1_CODE"),
                t2_code=_optional_code(value("t2_code"), row_number, "T2_CODE"),
                t3_code=_optional_code(value("t3_code"), row_number, "T3_CODE"),
                t4_code=_optional_code(value("t4_code"), row_number, "T4_CODE"),
                opening=_required_decimal(value("opening"), row_number, "OPENING"),
                debit=_required_decimal(value("debit"), row_number, "Debit"),
                credit=_required_decimal(value("credit"), row_number, "Credit"),
                closing=_required_decimal(value("closing"), row_number, "CLOSING"),
            )
    finally:
        workbook.close()


def _required_text(value: object, row_number: int, field: str) -> str:
    text = str(value).strip() if value is not None else ""
    if not text:
        raise LedgerParseError(f"row {row_number}: {field} is required")
    return text


def _required_year(value: object, row_number: int) -> int:
    parsed = _required_decimal(value, row_number, "year")
    if parsed != parsed.to_integral_value() or parsed <= 0:
        raise LedgerParseError(f"row {row_number}: year must be a positive integer")
    return int(parsed)


def _required_code(value: object, row_number: int, field: str) -> str:
    code = _optional_code(value, row_number, field)
    if code is None:
        raise LedgerParseError(f"row {row_number}: {field} is required")
    return code


def _optional_code(value: object, row_number: int, field: str) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        raise LedgerParseError(f"row {row_number}: {field} is not a valid code")
    if isinstance(value, str):
        code = value.strip()
        return code or None
    parsed = _required_decimal(value, row_number, field)
    if parsed != parsed.to_integral_value():
        raise LedgerParseError(f"row {row_number}: {field} must be an integer code")
    return str(int(parsed))


def _required_decimal(value: object, row_number: int, field: str) -> Decimal:
    if value is None or isinstance(value, bool):
        raise LedgerParseError(f"row {row_number}: {field} is required and numeric")
    candidate = str(value).strip().replace(",", "").replace("٬", "")
    try:
        parsed = Decimal(candidate)
    except (InvalidOperation, ValueError) as exc:
        raise LedgerParseError(f"row {row_number}: {field} is not numeric") from exc
    if not parsed.is_finite():
        raise LedgerParseError(f"row {row_number}: {field} must be finite")
    return parsed
