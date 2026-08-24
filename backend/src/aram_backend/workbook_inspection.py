from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


MAX_WORKSHEETS = 100
MAX_HEADER_COLUMNS = 256


@dataclass(frozen=True)
class WorkbookInspection:
    document_type: str
    sheet_names: tuple[str, ...]
    data_sheet: str | None
    header_row: int | None
    headers: tuple[str, ...]


class WorkbookInspectionError(ValueError):
    """Raised when an XLSX source does not match an accepted source profile."""


_TABULAR_PROFILES = {
    "general_ledger": {
        "sheet": "Source",
        "header_row": 2,
        "required_headers": frozenset(
            {
                "حساب",
                "ACCOUNT_NAME",
                "year",
                "نوع",
                "T1_CODE",
                "T2_CODE",
                "T3_CODE",
                "T4_CODE",
                "OPENING",
                "Debit",
                "Credit",
                "CLOSING",
            }
        ),
    },
    "mapping": {
        "sheet": "Data",
        "header_row": 1,
        "required_headers": frozenset(
            {
                "کد حساب",
                "نوع",
                "کد یادداشت",
                "شرح یادداشت",
                "فرمول یادداشت اصلی",
                "فرمول یادداشت فرعی",
            }
        ),
    },
}

_PRIMARY_STATEMENT_SHEETS = frozenset(
    {"سودوزيان", "وضعيت مالي", "حقوق مالكانه", "جريان هاي نقدي"}
)


def inspect_xlsx(content: bytes, *, document_type: str) -> WorkbookInspection:
    """Validate workbook structure without importing rows or evaluating formulas."""

    if not content:
        raise WorkbookInspectionError("workbook must not be empty")
    if document_type not in {*_TABULAR_PROFILES, "financial_statement"}:
        raise WorkbookInspectionError("document type has no XLSX inspection profile")

    try:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise WorkbookInspectionError("invalid or unreadable XLSX workbook") from exc

    try:
        sheet_names = tuple(workbook.sheetnames)
        if len(sheet_names) > MAX_WORKSHEETS:
            raise WorkbookInspectionError("workbook contains too many worksheets")

        if document_type == "financial_statement":
            missing = _PRIMARY_STATEMENT_SHEETS.difference(sheet_names)
            if missing:
                raise WorkbookInspectionError(
                    f"financial statement workbook is missing sheets: {sorted(missing)}"
                )
            return WorkbookInspection(
                document_type=document_type,
                sheet_names=sheet_names,
                data_sheet=None,
                header_row=None,
                headers=(),
            )

        profile = _TABULAR_PROFILES[document_type]
        sheet_name = profile["sheet"]
        if sheet_name not in workbook:
            raise WorkbookInspectionError(f"required worksheet is missing: {sheet_name}")

        worksheet = workbook[sheet_name]
        header_row = profile["header_row"]
        values = next(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                max_col=min(worksheet.max_column, MAX_HEADER_COLUMNS),
                values_only=True,
            ),
            (),
        )
        headers = tuple(str(value).strip() for value in values if value is not None)
        if len(headers) != len(set(headers)):
            raise WorkbookInspectionError("header row contains duplicate names")

        missing = profile["required_headers"].difference(headers)
        if missing:
            raise WorkbookInspectionError(f"required headers are missing: {sorted(missing)}")

        return WorkbookInspection(
            document_type=document_type,
            sheet_names=sheet_names,
            data_sheet=sheet_name,
            header_row=header_row,
            headers=headers,
        )
    finally:
        workbook.close()
