"""Explicit opt-in, nonofficial local analyst-review workbook export.

Every financial value is written as a text cell to preserve exact Decimal
scale and prevent spreadsheet formula injection. Existing paths are never
overwritten. This is an evidence/review pack, not an issued financial report.
"""
from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import unicodedata

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from aram_backend.coverage_preview import CoverageError
from aram_backend.desktop_five_statement_details import build_preview_details
from aram_backend.desktop_five_statement_figures import build_preview_figures
from aram_backend.desktop_five_statement_integrity import validate_desktop_pipeline_result
from aram_backend.five_statement_pipeline import FiveStatementPipelineResult

DISCLAIMER = (
    "غیررسمی — بسته بررسی کارشناسی — فاقد تأیید حسابداری و مجوز انتشار؛ "
    "به‌عنوان صورت مالی رسمی استفاده نشود"
)
MAX_REVIEW_WORKBOOK_BYTES = 40 * 1024 * 1024
MAX_EXPORT_FINDINGS = 10_000

SHEETS = (
    ("اطلاعات", "metadata"),
    ("خلاصه مبالغ", "summary"),
    ("ترازنامه", "ترازنامه"),
    ("سود و زیان", "سود و زیان"),
    ("حقوق مالکانه", "تغییرات حقوق مالکانه"),
    ("جریان نقد", "جریان وجوه نقد"),
    ("تراز آزمایشی", "تراز آزمایشی"),
    ("کنترل‌ها", "controls"),
)
DETAIL_HEADERS = (
    "صورت مالی", "شرح / حساب / فعالیت", "یادداشت / شناسه",
    "مواجهه / ردیف منبع", "افتتاحیه / منبع", "بدهکار / ورودی",
    "بستانکار / خروجی", "تغییر", "اختتامیه / ارائه",
    "مدرک / سیاست / ردگیری", "وضعیت",
)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
DISCLAIMER_FONT = Font(bold=True, color="9C0006")


def _safe_text(value: object) -> str:
    if not isinstance(value, str) or len(value) > 32_767:
        raise CoverageError("REVIEW_EXPORT_TEXT_INVALID")
    for char in value:
        category = unicodedata.category(char)
        codepoint = ord(char)
        if (category in {"Cc", "Cs"}
                or (category == "Cf" and codepoint not in {0x200C, 0x200D})):
            raise CoverageError("REVIEW_EXPORT_TEXT_CONTROL_INVALID")
    return value


def _cell(sheet, value: object, *, style: str = ""):
    text = _safe_text(value)
    cell = WriteOnlyCell(sheet, value=text)
    # openpyxl treats leading '=' as a formula by default. Force every
    # workbook field to a literal string, including labels/evidence IDs.
    cell.data_type = "s"
    if style == "header":
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    elif style == "disclaimer":
        cell.font = DISCLAIMER_FONT
    return cell


def _append(sheet, values, *, style: str = "") -> None:
    sheet.append([_cell(sheet, value, style=style) for value in values])


def _new_sheet(workbook: Workbook, title: str, metadata: str, headers: tuple[str, ...]):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A4"
    _append(sheet, (DISCLAIMER,), style="disclaimer")
    _append(sheet, (metadata,))
    sheet.append([])
    _append(sheet, headers, style="header")
    widths = (18, 30, 28, 28, 20, 20, 20, 20, 22, 55, 24)
    for index, width in enumerate(widths[:len(headers)], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    return sheet


def _metadata_rows(result: FiveStatementPipelineResult) -> tuple[tuple[str, str], ...]:
    state = (
        "قابل بررسی کارشناسی؛ غیررسمی"
        if result.eligible_for_analyst_review else "مسدود؛ صرفاً تشخیصی"
    )
    return (
        ("نوع خروجی", "UNAPPROVED_ANALYST_REVIEW_PACK"),
        ("شرکت", _safe_text(result.company_code)),
        ("سال مالی", str(result.fiscal_year)),
        ("واحد", _safe_text(result.cash_flow.unit)),
        ("ارز", _safe_text(result.cash_flow.currency)),
        ("وضعیت کلی", state),
        ("SHA256 فایل اکسل", result.source_sha256),
        ("SHA256 مراجع یادداشت", result.balance_sheet.note_references_sha256),
        ("SHA256 سیاست ترازنامه", result.balance_sheet.presentation_policies_sha256),
        ("SHA256 سیاست سود و زیان", result.income_statement.presentation_policies_sha256),
        ("SHA256 سیاست حقوق مالکانه", result.equity_changes.presentation_policies_sha256),
        ("SHA256 دفتر تراکنش نقد", result.cash_flow.transaction_ledger_sha256),
        ("SHA256 سیاست طبقه‌بندی نقد", result.cash_flow.classification_policies_sha256),
        ("SHA256 انتخاب یادداشت نقد", result.reconciliation.cash_note_selection_sha256),
        ("کامل‌بودن دفتر کل اثبات شده", "خیر"),
        ("سیاست حسابداری تأیید شده", "خیر"),
        ("بازبینی مستقل انجام شده", "خیر"),
        ("صورت مالی رسمی", "خیر"),
        ("مجوز انتشار", "خیر"),
        ("محدودیت", DISCLAIMER),
    )


def build_review_workbook_bytes(result: FiveStatementPipelineResult) -> bytes:
    """Build a bounded workbook entirely in memory; never touch source files."""
    validate_desktop_pipeline_result(result)
    figures = build_preview_figures(result)
    details = build_preview_details(result)
    findings = result.findings
    if not isinstance(findings, tuple) or len(findings) > MAX_EXPORT_FINDINGS:
        raise CoverageError("REVIEW_EXPORT_FINDINGS_LIMIT")

    workbook = Workbook(write_only=True)
    workbook.properties.title = "ARAM — UNOFFICIAL ANALYST REVIEW PACK"
    workbook.properties.subject = DISCLAIMER
    workbook.properties.creator = "ARAM Desktop Pilot"
    workbook.properties.description = DISCLAIMER

    info = _new_sheet(
        workbook, "اطلاعات", "مشخصات منبع و محدودیت‌های بسته بررسی",
        ("فیلد", "مقدار"),
    )
    for label, value in _metadata_rows(result):
        _append(info, (label, value))

    summary = _new_sheet(
        workbook, "خلاصه مبالغ", "۲۸ شاخص دقیق؛ مقادیر متنی و بدون گردکردن",
        ("صورت مالی", "شاخص", "مبلغ", "وضعیت"),
    )
    for figure in figures:
        _append(summary, (
            figure.statement, figure.label, figure.value, figure.state,
        ))

    for title, statement in SHEETS[2:7]:
        sheet = _new_sheet(
            workbook, title,
            "ریز اقلام منبع و ارائه — صرفاً پیش‌نویس غیررسمی",
            DETAIL_HEADERS,
        )
        for detail in details:
            if detail.statement == statement:
                _append(sheet, detail.values())

    controls = _new_sheet(
        workbook, "کنترل‌ها", "نتیجه کنترل‌های حسابی؛ بدون تأیید یا مجوز انتشار",
        ("کنترل", "مقدار", "وضعیت"),
    )
    reconciliation = result.reconciliation
    control_rows = (
        ("تطبیق ترازنامه", reconciliation.balance_sheet_difference,
         "معتبر" if reconciliation.balance_sheet_difference == "0" else "مغایرت"),
        ("اختلاف حقوق مالکانه", reconciliation.balance_sheet_equity_difference,
         "معتبر" if reconciliation.balance_sheet_equity_difference == "0" else "مغایرت"),
        ("اختلاف سود خالص", reconciliation.net_profit_equity_difference,
         "معتبر" if reconciliation.net_profit_equity_difference == "0" else "مغایرت"),
        ("اختلاف وجه نقد", reconciliation.balance_sheet_cash_difference,
         "معتبر" if reconciliation.balance_sheet_cash_difference == "0" else "مغایرت"),
        ("تطبیق سود و زیان", "بله" if reconciliation.income_statement_bridge_valid else "خیر",
         "کنترل حسابی"),
        ("تطبیق جریان نقد", "بله" if reconciliation.cash_flow_bridge_valid else "خیر",
         "کنترل حسابی"),
        ("گردش حقوق مالکانه", "بله" if reconciliation.equity_rollforward_valid else "خیر",
         "کنترل حسابی"),
        ("تراز آزمایشی", "بله" if reconciliation.trial_balance_arithmetic_valid else "خیر",
         "کنترل حسابی"),
        ("حساب‌های نقد با تراز آزمایشی", "بله" if reconciliation.cash_accounts_match_trial_balance else "خیر",
         "کنترل حسابی"),
        ("جمع حساب‌های نقد", "بله" if reconciliation.cash_account_totals_valid else "خیر",
         "کنترل حسابی"),
        ("منبع یکسان", "بله" if reconciliation.same_source_snapshot else "خیر",
         "کنترل منبع"),
        ("نتیجه کلی", "قابل بررسی" if result.eligible_for_analyst_review else "مسدود",
         "غیررسمی"),
    )
    for row in control_rows:
        _append(controls, row)
    _append(controls, ("مغایرت‌ها", "کد", "وضعیت"), style="header")
    for finding in findings:
        _append(controls, ("یافته", _safe_text(finding), "نیازمند بررسی"))

    buffer = BytesIO()
    try:
        workbook.save(buffer)
    finally:
        workbook.close()
    content = buffer.getvalue()
    if not content or len(content) > MAX_REVIEW_WORKBOOK_BYTES:
        raise CoverageError("REVIEW_EXPORT_SIZE_INVALID")
    return content


def write_review_workbook(
    path_value: str | os.PathLike[str],
    result: FiveStatementPipelineResult,
) -> Path:
    """Create a new local review workbook; fail rather than overwrite any path."""
    try:
        target = Path(path_value)
    except (TypeError, ValueError, OSError) as exc:
        raise CoverageError("REVIEW_EXPORT_PATH_INVALID") from exc
    if target.suffix.lower() != ".xlsx":
        raise CoverageError("REVIEW_EXPORT_EXTENSION_INVALID")
    if not target.parent.is_dir():
        raise CoverageError("REVIEW_EXPORT_PARENT_INVALID")
    if target.exists() or target.is_symlink():
        raise CoverageError("REVIEW_EXPORT_TARGET_EXISTS")

    content = build_review_workbook_bytes(result)
    flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL
    flags |= getattr(os, "O_BINARY", 0)
    try:
        descriptor = os.open(str(target), flags, 0o600)
    except FileExistsError as exc:
        raise CoverageError("REVIEW_EXPORT_TARGET_EXISTS") from exc
    except (OSError, ValueError) as exc:
        raise CoverageError("REVIEW_EXPORT_PATH_INVALID") from exc

    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
    except Exception as exc:
        try:
            target.unlink()
        except OSError:
            pass
        raise CoverageError("REVIEW_EXPORT_WRITE_FAILED") from exc
    return target
