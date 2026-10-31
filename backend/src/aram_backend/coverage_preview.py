"""Read-only preview of account-to-note Mapping coverage.

NEVER use the returned amounts as official financial statements. The workbook
can claim 'Approved', but a cell is not a verified reviewer identity. This
module performs no database writes, postings, imports or export of reports.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation, localcontext
from hashlib import sha256
from io import BytesIO
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from aram_backend.pilot_workbook import SHEET_HEADERS, preflight_pilot_workbook

MAX_INPUT_BYTES = 20 * 1024 * 1024
MAX_ZIP_ENTRIES = 500
MAX_EXPANDED_BYTES = 128 * 1024 * 1024


class CoverageError(ValueError):
    """A deterministic fail-closed coverage precondition violation."""

    def __init__(self, code: str) -> None:
        self.code = code
        super().__init__(code)


@dataclass(frozen=True)
class CoverageException:
    code: str
    account_code: str
    account_excel_row: int
    mapping_excel_rows: tuple[int, ...]


@dataclass(frozen=True)
class AccountTrace:
    account_code: str
    account_name: str
    account_excel_row: int
    source_row: int
    closing_exposure: str
    mapping_excel_row: int | None
    mapping_id: str | None
    mapping_count: int
    target_statement: str | None
    target_section: str | None
    target_note: str | None
    target_subnote: str | None
    included: bool


@dataclass(frozen=True)
class MappingTrace:
    account_code: str
    account_name: str
    account_excel_row: int
    source_row: int
    closing_exposure: str
    mapping_excel_row: int
    mapping_id: str
    target_statement: str
    target_section: str
    target_note: str
    target_subnote: str
    included: bool


@dataclass(frozen=True)
class NoteExposure:
    statement: str
    section: str
    note_number: str
    subnote_code: str
    closing_exposure: str
    account_codes: tuple[str, ...]


@dataclass(frozen=True)
class CoveragePreview:
    company_code: str
    fiscal_year: int
    source_sha256: str
    workbook_company_status: str
    account_count: int
    candidate_mapped_accounts: int
    blocked_accounts: int
    source_closing_exposure: str
    candidate_closing_exposure: str
    blocked_closing_exposure: str
    notes: tuple[NoteExposure, ...]
    accounts: tuple[AccountTrace, ...]
    mappings: tuple[MappingTrace, ...]
    exceptions: tuple[CoverageException, ...]
    kind: str = "UNAPPROVED_MAPPING_COVERAGE_PREVIEW"
    financial_statement: bool = False
    trusted_approvals: bool = False
    releasable: bool = False

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _guard_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_ENTRIES:
                raise CoverageError("WORKBOOK_ARCHIVE_LIMIT")
            expanded = 0
            for member in members:
                if member.file_size < 0 or member.file_size > MAX_EXPANDED_BYTES:
                    raise CoverageError("WORKBOOK_ARCHIVE_LIMIT")
                expanded += member.file_size
                if expanded > MAX_EXPANDED_BYTES:
                    raise CoverageError("WORKBOOK_ARCHIVE_LIMIT")
    except (BadZipFile, ValueError, OSError) as exc:
        if isinstance(exc, CoverageError):
            raise
        raise CoverageError("WORKBOOK_ARCHIVE_INVALID") from exc


def _read_rows(sheet: Any, headers: tuple[str, ...]) -> list[tuple[int, dict[str, Any]]]:
    rows = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        row = dict(zip(headers, values))
        if any(value is not None and str(value).strip() for value in row.values()):
            rows.append((excel_row, row))
    return rows


def _literal_code(value: Any) -> str:
    if not isinstance(value, str) or not value.strip() or value.lstrip().startswith("="):
        raise CoverageError("INVALID_ACCOUNT_OR_NOTE_CODE")
    return value.strip()


def _closing_decimal(value: Any) -> Decimal:
    if value is None or isinstance(value, bool):
        raise CoverageError("INVALID_CLOSING")
    try:
        amount = Decimal(str(value).strip())
    except (TypeError, ValueError, InvalidOperation) as exc:
        raise CoverageError("INVALID_CLOSING") from exc
    if not amount.is_finite() or (amount != 0 and abs(amount.adjusted()) > 70):
        raise CoverageError("UNSUPPORTED_CLOSING_MAGNITUDE")
    return amount


def _assert_visible(sheet: Any) -> None:
    if sheet.sheet_state != "visible":
        raise CoverageError("HIDDEN_INPUT")
    if any(item.hidden for item in sheet.row_dimensions.values()):
        raise CoverageError("HIDDEN_INPUT")
    if any(item.hidden for item in sheet.column_dimensions.values()):
        raise CoverageError("HIDDEN_INPUT")
    if sheet.merged_cells.ranges:
        raise CoverageError("MERGED_INPUT")


def preview_mapping_coverage(
    content: bytes, *, company_code: str, fiscal_year: int
) -> CoveragePreview:
    """Preview one source account against one mapping per target statement.

    One account may feed multiple distinct statements (for example BS and EQ),
    but may not feed multiple notes within the same statement. Each target is
    independently file-marked Approved+Valid; those flags remain untrusted.
    Source totals count each account once, while note exposures are grouped by
    every approved mapping. No financial statement or approval is created.
    """
    if not isinstance(company_code, str) or not company_code.strip():
        raise CoverageError("COMPANY_SCOPE_REQUIRED")
    if type(fiscal_year) is not int or fiscal_year <= 0:
        raise CoverageError("FISCAL_YEAR_SCOPE_REQUIRED")
    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_INPUT_BYTES:
        raise CoverageError("WORKBOOK_SIZE_INVALID")
    _guard_archive(content)

    try:
        preflight = preflight_pilot_workbook(content)
    except Exception as exc:
        raise CoverageError("WORKBOOK_INVALID") from exc
    if not preflight.valid:
        raise CoverageError("WORKBOOK_INVALID")

    scope = (company_code.strip(), fiscal_year)

    def scoped(row: dict[str, Any]) -> bool:
        return (
            isinstance(row.get("company_code"), str)
            and row["company_code"].strip() == scope[0]
            and type(row.get("fiscal_year")) is int
            and row["fiscal_year"] == scope[1]
        )

    try:
        book = load_workbook(
            BytesIO(content), read_only=False, data_only=False, keep_links=False,
        )
    except Exception as exc:
        raise CoverageError("WORKBOOK_INVALID") from exc

    try:
        for sheet in book.worksheets:
            _assert_visible(sheet)
        for name, expected in SHEET_HEADERS.items():
            if name not in book:
                raise CoverageError("WORKBOOK_INVALID")
            if tuple(cell.value for cell in book[name][1]) != expected:
                raise CoverageError("WORKBOOK_INVALID")

        company_rows = [
            row for _, row in _read_rows(book["Company"], SHEET_HEADERS["Company"])
            if scoped(row)
        ]
        if len(company_rows) != 1:
            raise CoverageError("COMPANY_SCOPE_NOT_FOUND")
        status = str(company_rows[0]["status"])

        accounts: dict[str, tuple[int, dict[str, Any], Decimal]] = {}
        for excel_row, row in _read_rows(book["Accounts"], SHEET_HEADERS["Accounts"]):
            if not scoped(row):
                continue
            code = _literal_code(row["account_full_code"])
            if code in accounts:
                raise CoverageError("DUPLICATE_ACCOUNT")
            accounts[code] = (excel_row, row, _closing_decimal(row["closing"]))
        if not accounts:
            raise CoverageError("NO_ACCOUNTS_IN_SCOPE")

        note_keys = {
            (
                str(row["statement"]).strip(),
                str(row["statement_section"]).strip(),
                _literal_code(row["note_number"]),
                row["subnote_code"].strip() if isinstance(row["subnote_code"], str) else "",
            )
            for _, row in _read_rows(book["Notes"], SHEET_HEADERS["Notes"])
            if scoped(row)
        }

        mappings_by_account: dict[str, list[tuple[int, dict[str, Any]]]] = defaultdict(list)
        for excel_row, row in _read_rows(book["Mapping"], SHEET_HEADERS["Mapping"]):
            if not scoped(row):
                continue
            code = _literal_code(row["account_full_code"])
            if code not in accounts:
                raise CoverageError("UNKNOWN_MAPPING_ACCOUNT")
            mappings_by_account[code].append((excel_row, row))

        exceptions: list[CoverageException] = []
        account_traces: list[AccountTrace] = []
        mapping_traces: list[MappingTrace] = []
        grouped: dict[tuple[str, str, str, str], list[tuple[str, Decimal]]] = defaultdict(list)
        eligible_amounts: list[Decimal] = []
        blocked_amounts: list[Decimal] = []
        all_amounts: list[Decimal] = []

        for code in sorted(accounts):
            excel_row, source, closing = accounts[code]
            all_amounts.append(closing)
            matches = mappings_by_account.get(code, [])
            parsed: list[tuple[int, str, tuple[str, str, str, str], dict[str, Any]]] = []
            by_statement: dict[str, list[tuple[int, str, tuple[str, str, str, str], dict[str, Any]]]] = defaultdict(list)

            for mapping_row, mapping in matches:
                mapping_id = _literal_code(mapping["mapping_id"])
                target = (
                    str(mapping["target_statement"]).strip(),
                    str(mapping["statement_section"]).strip(),
                    _literal_code(mapping["note_number"]),
                    mapping["subnote_code"].strip()
                    if isinstance(mapping["subnote_code"], str) else "",
                )
                if target not in note_keys:
                    raise CoverageError("UNKNOWN_MAPPING_NOTE")
                item = (mapping_row, mapping_id, target, mapping)
                parsed.append(item)
                by_statement[target[0]].append(item)

            included_rows: list[tuple[int, str, tuple[str, str, str, str], dict[str, Any]]] = []
            if not matches:
                exceptions.append(CoverageException("UNMAPPED_ACCOUNT", code, excel_row, ()))
            else:
                for statement in sorted(by_statement):
                    statement_rows = by_statement[statement]
                    if len(statement_rows) > 1:
                        exceptions.append(CoverageException(
                            "AMBIGUOUS_ACCOUNT_MAPPING", code, excel_row,
                            tuple(sorted(item[0] for item in statement_rows)),
                        ))
                        continue
                    mapping_row, mapping_id, target, mapping = statement_rows[0]
                    if (mapping["approval_status"] != "Approved"
                            or mapping["validation_status"] != "Valid"):
                        exceptions.append(CoverageException(
                            "MAPPING_NOT_FILE_MARKED_APPROVED", code, excel_row,
                            (mapping_row,),
                        ))
                        continue
                    included_rows.append((mapping_row, mapping_id, target, mapping))
                    grouped[target].append((code, closing))

            candidate = bool(included_rows)
            if candidate:
                eligible_amounts.append(closing)
            else:
                blocked_amounts.append(closing)

            for mapping_row, mapping_id, target, mapping in parsed:
                mapping_traces.append(MappingTrace(
                    account_code=code, account_name=str(source["account_name"]),
                    account_excel_row=excel_row, source_row=int(source["source_row"]),
                    closing_exposure=str(closing), mapping_excel_row=mapping_row,
                    mapping_id=mapping_id, target_statement=target[0],
                    target_section=target[1], target_note=target[2],
                    target_subnote=target[3],
                    included=any(row[0] == mapping_row for row in included_rows),
                ))

            single = parsed[0] if len(parsed) == 1 else None
            single_included = bool(single and any(row[0] == single[0] for row in included_rows))
            account_traces.append(AccountTrace(
                account_code=code, account_name=str(source["account_name"]),
                account_excel_row=excel_row, source_row=int(source["source_row"]),
                closing_exposure=str(closing),
                mapping_excel_row=single[0] if single else None,
                mapping_id=single[1] if single else None,
                mapping_count=len(matches),
                target_statement=single[2][0] if single and single_included else None,
                target_section=single[2][1] if single and single_included else None,
                target_note=single[2][2] if single and single_included else None,
                target_subnote=single[2][3] if single and single_included else None,
                included=candidate,
            ))

        positive_digits = max(
            (amount.adjusted() + 1 for amount in all_amounts if amount),
            default=1,
        )
        fractional_digits = max(0, *(-amount.as_tuple().exponent for amount in all_amounts))
        precision = max(
            100, positive_digits + fractional_digits + len(str(len(all_amounts))) + 4,
        )
        with localcontext() as ctx:
            ctx.prec = precision
            source_total = sum(all_amounts, Decimal(0))
            eligible_total = sum(eligible_amounts, Decimal(0))
            blocked_total = sum(blocked_amounts, Decimal(0))
            if source_total != eligible_total + blocked_total:
                raise CoverageError("COVERAGE_RECONCILIATION_FAILED")
            note_results = tuple(
                NoteExposure(
                    statement=key[0], section=key[1],
                    note_number=key[2], subnote_code=key[3],
                    closing_exposure=str(sum(
                        (amount for _, amount in grouped[key]), Decimal(0),
                    )),
                    account_codes=tuple(code for code, _ in grouped[key]),
                )
                for key in sorted(grouped)
            )

        return CoveragePreview(
            company_code=scope[0], fiscal_year=fiscal_year,
            source_sha256=sha256(content).hexdigest(),
            workbook_company_status=status, account_count=len(accounts),
            candidate_mapped_accounts=len(eligible_amounts),
            blocked_accounts=len(blocked_amounts),
            source_closing_exposure=str(source_total),
            candidate_closing_exposure=str(eligible_total),
            blocked_closing_exposure=str(blocked_total),
            notes=note_results, accounts=tuple(account_traces),
            mappings=tuple(sorted(
                mapping_traces,
                key=lambda item: (item.account_code, item.mapping_excel_row),
            )),
            exceptions=tuple(sorted(
                exceptions, key=lambda entry: (
                    entry.account_code, entry.code, entry.mapping_excel_rows,
                ),
            )),
        )
    finally:
        book.close()
