"""Direct-method cash flow candidate from paired, classified cash journal lines.

This module deliberately requires transaction-level cash/counterpart postings.
Aggregate account debit/credit totals alone cannot identify operating, investing
or financing activity and must never be guessed into a cash flow statement.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
import re
from typing import Any, Sequence

from openpyxl import load_workbook

from aram_backend.coverage_preview import (
    CoverageError, MAX_INPUT_BYTES, _assert_visible, _guard_archive,
)
from aram_backend.pilot_workbook import SHEET_HEADERS, preflight_pilot_workbook
from aram_backend.reconciliation_review import _money, _sum_exact

MAX_CASH_ACCOUNTS = 5_000
MAX_COUNTERPART_POLICIES = 20_000
MAX_CASH_TRANSACTIONS = 30_000
ACTIVITIES = frozenset({"OPERATING", "INVESTING", "FINANCING"})


@dataclass(frozen=True)
class CashAccountPolicy:
    account_code: str
    evidence_ref: str


@dataclass(frozen=True)
class CashCounterpartPolicy:
    account_code: str
    activity: str
    evidence_ref: str


@dataclass(frozen=True)
class CashFlowEntry:
    company_code: str
    fiscal_year: int
    line_id: str
    transaction_id: str
    posting_date: str
    cash_account_code: str
    counterpart_account_code: str
    cash_debit: str
    cash_credit: str
    counterpart_debit: str
    counterpart_credit: str
    evidence_ref: str


@dataclass(frozen=True)
class CashAccountReconciliation:
    account_code: str
    account_name: str
    source_row: int
    opening_balance: str
    source_debit_turnover: str
    source_credit_turnover: str
    closing_balance: str
    transaction_debit_turnover: str
    transaction_credit_turnover: str
    debit_difference: str
    credit_difference: str
    rollforward_difference: str
    policy_evidence_ref: str


@dataclass(frozen=True)
class CashFlowLine:
    line_id: str
    transaction_id: str
    posting_date: str
    cash_account_code: str
    counterpart_account_code: str
    activity: str
    cash_inflow: str
    cash_outflow: str
    source_evidence_ref: str
    classification_evidence_ref: str | None


@dataclass(frozen=True)
class DirectCashFlowCandidate:
    company_code: str
    fiscal_year: int
    currency: str
    unit: str
    source_sha256: str
    transaction_ledger_sha256: str
    classification_policies_sha256: str
    cash_account_count: int
    transaction_count: int
    opening_cash_and_equivalents: str
    operating_inflows: str
    operating_outflows: str
    operating_net_cash_flow: str
    investing_inflows: str
    investing_outflows: str
    investing_net_cash_flow: str
    financing_inflows: str
    financing_outflows: str
    financing_net_cash_flow: str
    internal_transfer_inflows: str
    internal_transfer_outflows: str
    net_change_in_cash: str
    closing_cash_and_equivalents: str
    rollforward_difference: str
    account_reconciliations: tuple[CashAccountReconciliation, ...]
    cash_flow_lines: tuple[CashFlowLine, ...]
    findings: tuple[str, ...]
    eligible_for_analyst_review: bool
    kind: str = "UNAPPROVED_DIRECT_CASH_FLOW_CANDIDATE"
    cash_account_scope_owner_approved: bool = False
    transaction_ledger_completeness_verified: bool = False
    accounting_policy_authorized: bool = False
    official_financial_statement: bool = False
    independently_reviewed: bool = False
    export_authorized: bool = False
    releasable: bool = False

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _text(value: object, label: str, *, max_length: int = 256) -> str:
    if (not isinstance(value, str) or not value or len(value) > max_length
            or value != value.strip()
            or any(ord(char) < 32 or ord(char) == 127 for char in value)):
        raise CoverageError("INVALID_CASH_FLOW_" + label)
    return value


def _date(value: object, fiscal_year: int) -> str:
    result = _text(value, "POSTING_DATE", max_length=10)
    if (not re.fullmatch(r"\d{4}-\d{2}-\d{2}", result)
            or int(result[:4]) != fiscal_year
            or not 1 <= int(result[5:7]) <= 12
            or not 1 <= int(result[8:10]) <= 31):
        raise CoverageError("INVALID_CASH_FLOW_POSTING_DATE")
    # Year is the explicitly selected source fiscal-year label; no Gregorian
    # or Jalali conversion is attempted here.
    return result


def _policies(
    cash_accounts: Sequence[CashAccountPolicy],
    counterparts: Sequence[CashCounterpartPolicy],
) -> tuple[dict[str, CashAccountPolicy], dict[str, CashCounterpartPolicy], str]:
    if (not isinstance(cash_accounts, (list, tuple))
            or not 0 < len(cash_accounts) <= MAX_CASH_ACCOUNTS):
        raise CoverageError("INVALID_CASH_ACCOUNT_POLICIES")
    if (not isinstance(counterparts, (list, tuple))
            or len(counterparts) > MAX_COUNTERPART_POLICIES):
        raise CoverageError("INVALID_CASH_COUNTERPART_POLICIES")
    cash: dict[str, CashAccountPolicy] = {}
    other: dict[str, CashCounterpartPolicy] = {}
    for policy in cash_accounts:
        if not isinstance(policy, CashAccountPolicy):
            raise CoverageError("INVALID_CASH_ACCOUNT_POLICY")
        code = _text(policy.account_code, "ACCOUNT_CODE")
        evidence = _text(policy.evidence_ref, "ACCOUNT_EVIDENCE")
        if code in cash:
            raise CoverageError("DUPLICATE_CASH_ACCOUNT_POLICY")
        cash[code] = policy
    for policy in counterparts:
        if not isinstance(policy, CashCounterpartPolicy):
            raise CoverageError("INVALID_CASH_COUNTERPART_POLICY")
        code = _text(policy.account_code, "COUNTERPART_CODE")
        evidence = _text(policy.evidence_ref, "COUNTERPART_EVIDENCE")
        if code in cash:
            raise CoverageError("CASH_ACCOUNT_CANNOT_BE_COUNTERPART_POLICY")
        if policy.activity not in ACTIVITIES:
            raise CoverageError("INVALID_CASH_FLOW_ACTIVITY")
        if code in other:
            raise CoverageError("DUPLICATE_CASH_COUNTERPART_POLICY")
        other[code] = policy
    canonical = {
        "cash_accounts": [
            asdict(cash[key]) for key in sorted(cash)
        ],
        "counterparts": [
            asdict(other[key]) for key in sorted(other)
        ],
    }
    digest = sha256(json.dumps(
        canonical, ensure_ascii=True, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    return cash, other, digest


def build_direct_cash_flow_candidate(
    workbook_bytes: bytes, *, company_code: str, fiscal_year: int,
    currency: str, unit: str,
    cash_accounts: Sequence[CashAccountPolicy],
    counterpart_policies: Sequence[CashCounterpartPolicy],
    transactions: Sequence[CashFlowEntry],
) -> DirectCashFlowCandidate:
    """Build a direct-method analyst preview using paired cash journal entries.

    Every row represents exactly one cash posting paired with its opposite
    counterpart posting. If a voucher has multiple counterpart lines, provide
    one separately identified pair per line. Cash-to-cash transfers are
    recognized only when both cash accounts are explicitly registered and are
    excluded from operating/investing/financing totals.
    """
    company_code = _text(company_code, "COMPANY_SCOPE", max_length=64)
    if type(fiscal_year) is not int or fiscal_year <= 0:
        raise CoverageError("CASH_FLOW_YEAR_SCOPE_INVALID")
    currency = _text(currency, "CURRENCY", max_length=16)
    unit = _text(unit, "UNIT", max_length=32)
    cash_policy, counterpart_policy, policy_hash = _policies(
        cash_accounts, counterpart_policies,
    )
    if (not isinstance(transactions, (list, tuple))
            or len(transactions) > MAX_CASH_TRANSACTIONS):
        raise CoverageError("INVALID_CASH_FLOW_TRANSACTIONS")
    if (not isinstance(workbook_bytes, bytes)
            or not 0 < len(workbook_bytes) <= MAX_INPUT_BYTES):
        raise CoverageError("CASH_FLOW_WORKBOOK_SIZE_INVALID")
    _guard_archive(workbook_bytes)
    try:
        preflight = preflight_pilot_workbook(workbook_bytes)
    except Exception as exc:
        raise CoverageError("CASH_FLOW_WORKBOOK_INVALID") from exc
    if not preflight.valid:
        raise CoverageError("CASH_FLOW_WORKBOOK_INVALID")
    try:
        book = load_workbook(
            BytesIO(workbook_bytes), read_only=False,
            data_only=False, keep_links=False,
        )
    except Exception as exc:
        raise CoverageError("CASH_FLOW_WORKBOOK_INVALID") from exc

    try:
        for sheet in book.worksheets:
            _assert_visible(sheet)
        for name, headers in SHEET_HEADERS.items():
            if name not in book or tuple(cell.value for cell in book[name][1]) != headers:
                raise CoverageError("CASH_FLOW_WORKBOOK_INVALID")

        scope_company = company_code.strip()
        companies = []
        for values in book["Company"].iter_rows(min_row=2, values_only=True):
            row = dict(zip(SHEET_HEADERS["Company"], values))
            if (isinstance(row.get("company_code"), str)
                    and row["company_code"].strip() == scope_company
                    and type(row.get("fiscal_year")) is int
                    and row["fiscal_year"] == fiscal_year):
                companies.append(row)
        if len(companies) != 1:
            raise CoverageError("CASH_FLOW_COMPANY_NOT_FOUND")
        company = companies[0]
        if company.get("currency") != currency or company.get("unit") != unit:
            raise CoverageError("CASH_FLOW_CURRENCY_UNIT_MISMATCH")

        source_accounts: dict[str, tuple[int, dict[str, Any]]] = {}
        for row_number, values in enumerate(
            book["Accounts"].iter_rows(min_row=2, values_only=True), 2,
        ):
            row = dict(zip(SHEET_HEADERS["Accounts"], values))
            if not any(value is not None and str(value).strip() for value in row.values()):
                continue
            if (isinstance(row.get("company_code"), str)
                    and row["company_code"].strip() == scope_company
                    and type(row.get("fiscal_year")) is int
                    and row["fiscal_year"] == fiscal_year):
                code = row.get("account_full_code")
                if not isinstance(code, str) or not code.strip():
                    raise CoverageError("CASH_FLOW_ACCOUNT_CODE_INVALID")
                code = code.strip()
                if code in source_accounts:
                    raise CoverageError("CASH_FLOW_DUPLICATE_SOURCE_ACCOUNT")
                source_accounts[code] = (row_number, row)

        if set(cash_policy) - set(source_accounts):
            raise CoverageError("CASH_ACCOUNT_NOT_FOUND_IN_SOURCE")

        source_balances: dict[str, tuple[Decimal, Decimal, Decimal, Decimal, int, str]] = {}
        for code, policy in cash_policy.items():
            row_number, row = source_accounts[code]
            opening, debit, credit, closing = (
                _money(str(row[field]))
                for field in ("opening", "debit", "credit", "closing")
            )
            if debit < 0 or credit < 0:
                raise CoverageError("CASH_ACCOUNT_NEGATIVE_TURNOVER")
            row_difference = _sum_exact([opening, debit, -credit, -closing])
            if row_difference != 0:
                raise CoverageError("CASH_ACCOUNT_ROLLFORWARD_INVALID")
            if type(row.get("source_row")) is not int or row["source_row"] <= 0:
                raise CoverageError("CASH_ACCOUNT_SOURCE_ROW_INVALID")
            source_balances[code] = (
                opening, debit, credit, closing, row["source_row"],
                str(row["account_name"]),
            )

        line_ids: set[str] = set()
        used_counterparts: set[str] = set()
        transaction_turnovers: dict[str, tuple[list[Decimal], list[Decimal]]] = {
            code: ([], []) for code in cash_policy
        }
        activity_totals: dict[str, dict[str, list[Decimal]]] = {
            activity: {"inflows": [], "outflows": []}
            for activity in sorted(ACTIVITIES)
        }
        transfer_inflows: list[Decimal] = []
        transfer_outflows: list[Decimal] = []
        output_lines: list[CashFlowLine] = []
        canonical_transactions: list[dict[str, Any]] = []
        findings: list[str] = []

        for entry in transactions:
            if not isinstance(entry, CashFlowEntry):
                raise CoverageError("INVALID_CASH_FLOW_TRANSACTION")
            company_value = _text(entry.company_code, "TRANSACTION_COMPANY")
            if company_value != scope_company or type(entry.fiscal_year) is not int or entry.fiscal_year != fiscal_year:
                raise CoverageError("CASH_FLOW_TRANSACTION_SCOPE_MISMATCH")
            line_id = _text(entry.line_id, "LINE_ID")
            transaction_id = _text(entry.transaction_id, "TRANSACTION_ID")
            posting_date = _date(entry.posting_date, fiscal_year)
            cash_code = _text(entry.cash_account_code, "CASH_ACCOUNT_CODE")
            counterpart_code = _text(entry.counterpart_account_code, "COUNTERPART_ACCOUNT_CODE")
            evidence = _text(entry.evidence_ref, "TRANSACTION_EVIDENCE")
            if line_id in line_ids:
                raise CoverageError("DUPLICATE_CASH_FLOW_LINE_ID")
            line_ids.add(line_id)
            if cash_code not in cash_policy:
                raise CoverageError("UNREGISTERED_CASH_ACCOUNT_IN_TRANSACTION")
            if counterpart_code == cash_code or counterpart_code not in source_accounts:
                raise CoverageError("INVALID_CASH_COUNTERPART_ACCOUNT")

            cash_debit, cash_credit, counterpart_debit, counterpart_credit = (
                _money(value) for value in (
                    entry.cash_debit, entry.cash_credit,
                    entry.counterpart_debit, entry.counterpart_credit,
                )
            )
            if min(cash_debit, cash_credit, counterpart_debit, counterpart_credit) < 0:
                raise CoverageError("NEGATIVE_CASH_FLOW_POSTING")
            debit_side = cash_debit > 0 and cash_credit == 0
            credit_side = cash_credit > 0 and cash_debit == 0
            if not (debit_side or credit_side):
                raise CoverageError("CASH_FLOW_ENTRY_MUST_HAVE_ONE_CASH_SIDE")
            if debit_side:
                if counterpart_credit != cash_debit or counterpart_debit != 0:
                    raise CoverageError("CASH_FLOW_COUNTERPART_NOT_BALANCED")
            elif counterpart_debit != cash_credit or counterpart_credit != 0:
                raise CoverageError("CASH_FLOW_COUNTERPART_NOT_BALANCED")

            transaction_turnovers[cash_code][0].append(cash_debit)
            transaction_turnovers[cash_code][1].append(cash_credit)
            if counterpart_code in cash_policy:
                transaction_turnovers[counterpart_code][0].append(counterpart_debit)
                transaction_turnovers[counterpart_code][1].append(counterpart_credit)
                activity = "CASH_TRANSFER"
                classification_evidence = cash_policy[counterpart_code].evidence_ref
                transfer_inflows.extend([cash_debit, counterpart_debit])
                transfer_outflows.extend([cash_credit, counterpart_credit])
            else:
                policy = counterpart_policy.get(counterpart_code)
                if policy is None:
                    activity = "UNCLASSIFIED"
                    classification_evidence = None
                    findings.append("MISSING_COUNTERPART_ACTIVITY_POLICY")
                else:
                    used_counterparts.add(counterpart_code)
                    activity = policy.activity
                    classification_evidence = policy.evidence_ref
                    if cash_debit:
                        activity_totals[activity]["inflows"].append(cash_debit)
                    else:
                        activity_totals[activity]["outflows"].append(cash_credit)

            output_lines.append(CashFlowLine(
                line_id=line_id, transaction_id=transaction_id,
                posting_date=posting_date, cash_account_code=cash_code,
                counterpart_account_code=counterpart_code, activity=activity,
                cash_inflow=str(cash_debit), cash_outflow=str(cash_credit),
                source_evidence_ref=evidence,
                classification_evidence_ref=classification_evidence,
            ))
            canonical_transactions.append(asdict(entry))

        if set(counterpart_policy) - used_counterparts:
            findings.append("ORPHAN_COUNTERPART_ACTIVITY_POLICY")

        reconciliations: list[CashAccountReconciliation] = []
        opening_values: list[Decimal] = []
        closing_values: list[Decimal] = []
        for code in sorted(cash_policy):
            opening, source_debit, source_credit, closing, source_row, name = source_balances[code]
            txn_debit = _sum_exact(transaction_turnovers[code][0])
            txn_credit = _sum_exact(transaction_turnovers[code][1])
            debit_diff = _sum_exact([txn_debit, -source_debit])
            credit_diff = _sum_exact([txn_credit, -source_credit])
            rollforward = _sum_exact([opening, source_debit, -source_credit, -closing])
            if debit_diff != 0:
                findings.append("CASH_ACCOUNT_DEBIT_TURNOVER_MISMATCH")
            if credit_diff != 0:
                findings.append("CASH_ACCOUNT_CREDIT_TURNOVER_MISMATCH")
            if rollforward != 0:
                findings.append("CASH_ACCOUNT_ROLLFORWARD_MISMATCH")
            reconciliations.append(CashAccountReconciliation(
                account_code=code, account_name=name, source_row=source_row,
                opening_balance=str(opening),
                source_debit_turnover=str(source_debit),
                source_credit_turnover=str(source_credit),
                closing_balance=str(closing),
                transaction_debit_turnover=str(txn_debit),
                transaction_credit_turnover=str(txn_credit),
                debit_difference=str(debit_diff), credit_difference=str(credit_diff),
                rollforward_difference=str(rollforward),
                policy_evidence_ref=cash_policy[code].evidence_ref,
            ))
            opening_values.append(opening)
            closing_values.append(closing)

        totals = {
            activity: {
                "inflows": _sum_exact(activity_totals[activity]["inflows"]),
                "outflows": _sum_exact(activity_totals[activity]["outflows"]),
            }
            for activity in sorted(ACTIVITIES)
        }
        net_by_activity = {
            activity: _sum_exact([
                totals[activity]["inflows"], -totals[activity]["outflows"],
            ])
            for activity in sorted(ACTIVITIES)
        }
        opening_total = _sum_exact(opening_values)
        closing_total = _sum_exact(closing_values)
        net_change = _sum_exact(list(net_by_activity.values()))
        rollforward_difference = _sum_exact([
            opening_total, net_change, -closing_total,
        ])
        if rollforward_difference != 0:
            findings.append("CASH_FLOW_CASH_BALANCE_ROLLFORWARD_MISMATCH")

        transaction_hash = sha256(json.dumps(
            sorted(canonical_transactions, key=lambda row: row["line_id"]),
            ensure_ascii=True, sort_keys=True, separators=(",", ":"),
        ).encode("utf-8")).hexdigest()
        return DirectCashFlowCandidate(
            company_code=scope_company, fiscal_year=fiscal_year,
            currency=currency, unit=unit,
            source_sha256=sha256(workbook_bytes).hexdigest(),
            transaction_ledger_sha256=transaction_hash,
            classification_policies_sha256=policy_hash,
            cash_account_count=len(cash_policy),
            transaction_count=len(transactions),
            opening_cash_and_equivalents=str(opening_total),
            operating_inflows=str(totals["OPERATING"]["inflows"]),
            operating_outflows=str(totals["OPERATING"]["outflows"]),
            operating_net_cash_flow=str(net_by_activity["OPERATING"]),
            investing_inflows=str(totals["INVESTING"]["inflows"]),
            investing_outflows=str(totals["INVESTING"]["outflows"]),
            investing_net_cash_flow=str(net_by_activity["INVESTING"]),
            financing_inflows=str(totals["FINANCING"]["inflows"]),
            financing_outflows=str(totals["FINANCING"]["outflows"]),
            financing_net_cash_flow=str(net_by_activity["FINANCING"]),
            internal_transfer_inflows=str(_sum_exact(transfer_inflows)),
            internal_transfer_outflows=str(_sum_exact(transfer_outflows)),
            net_change_in_cash=str(net_change),
            closing_cash_and_equivalents=str(closing_total),
            rollforward_difference=str(rollforward_difference),
            account_reconciliations=tuple(reconciliations),
            cash_flow_lines=tuple(sorted(
                output_lines, key=lambda line: (line.posting_date, line.line_id),
            )),
            findings=tuple(sorted(set(findings))),
            eligible_for_analyst_review=not findings and bool(cash_policy),
        )
    finally:
        book.close()
