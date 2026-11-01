"""Governed, isolated mapping decisions for a single company and fiscal year.

No approval is inferred from an imported workbook. All transitions are explicit
and append-only audit entries; caller owns durable storage and transactionality.
"""
from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass, replace
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from aram_backend.pilot_workbook import SHEET_HEADERS, preflight_pilot_workbook


class MappingWorkflowError(ValueError):
    pass


@dataclass(frozen=True)
class MappingDecision:
    mapping_id: str
    company_code: str
    fiscal_year: int
    account_full_code: str
    target_statement: str
    statement_section: str
    note_number: str
    subnote_code: str | None
    version: int
    status: str
    reviewer: str | None
    source: str
    review_date: str | None = None


@dataclass(frozen=True)
class MappingAuditEvent:
    mapping_id: str
    company_code: str
    fiscal_year: int
    previous_status: str | None
    new_status: str
    actor: str
    timestamp: str
    previous_version: int | None
    new_version: int


class MappingWorkspace:
    """In-memory transaction boundary; never crosses the selected company/year."""

    def __init__(self, company_code: str, fiscal_year: int):
        if not isinstance(company_code, str) or not company_code.strip():
            raise MappingWorkflowError("company_code is required")
        if isinstance(fiscal_year, bool) or not isinstance(fiscal_year, int) or fiscal_year <= 0:
            raise MappingWorkflowError("fiscal_year must be a positive integer")
        self.company_code = company_code.strip()
        self.fiscal_year = fiscal_year
        self._decisions: dict[str, MappingDecision] = {}
        self._audit: list[MappingAuditEvent] = []

    @property
    def decisions(self) -> tuple[MappingDecision, ...]:
        return tuple(self._decisions[key] for key in sorted(self._decisions))

    @property
    def audit(self) -> tuple[MappingAuditEvent, ...]:
        return tuple(self._audit)

    def _check_scope(self, decision: MappingDecision) -> None:
        if (decision.company_code, decision.fiscal_year) != (self.company_code, self.fiscal_year):
            raise MappingWorkflowError("company/year isolation violation")

    def _record(self, old: MappingDecision | None, new: MappingDecision, actor: str) -> None:
        self._decisions[new.mapping_id] = new
        self._audit.append(MappingAuditEvent(
            new.mapping_id, new.company_code, new.fiscal_year,
            old.status if old else None, new.status, actor,
            datetime.now(UTC).isoformat(),
            old.version if old else None, new.version,
        ))

    def propose(self, decision: MappingDecision, *, actor: str) -> MappingDecision:
        self._check_scope(decision)
        if not actor.strip() or not decision.mapping_id.strip() or not decision.account_full_code.strip():
            raise MappingWorkflowError("actor, mapping_id and account are required")
        if (decision.status != "Draft" or decision.version != 1
                or decision.reviewer is not None or decision.review_date is not None):
            raise MappingWorkflowError("new decisions must begin as unreviewed Draft v1")
        if decision.mapping_id in self._decisions:
            raise MappingWorkflowError("mapping_id already exists; use revise")
        if not decision.note_number or not decision.target_statement or not decision.statement_section:
            raise MappingWorkflowError("target statement, section and note are required")
        self._record(None, decision, actor)
        return decision

    def transition(self, mapping_id: str, status: str, *, actor: str) -> MappingDecision:
        old = self._decisions.get(mapping_id)
        if old is None:
            raise MappingWorkflowError("unknown mapping_id")
        allowed = {"Draft": "Reviewed", "Reviewed": "Approved"}
        if allowed.get(old.status) != status:
            raise MappingWorkflowError("invalid mapping lifecycle transition")
        if not actor.strip():
            raise MappingWorkflowError("reviewer identity required")
        if status == "Reviewed" and any(
            event.mapping_id == mapping_id and event.new_status == "Draft"
            and event.actor == actor for event in self._audit
        ):
            raise MappingWorkflowError("maker cannot review own mapping")
        if status == "Approved" and (actor == old.reviewer or any(
            event.mapping_id == mapping_id and event.new_status == "Draft"
            and event.actor == actor for event in self._audit
        )):
            raise MappingWorkflowError("approval requires independent maker and reviewer")
        new = replace(old, status=status, reviewer=actor,
                      review_date=datetime.now(UTC).date().isoformat(),
                      version=old.version + 1)
        self._record(old, new, actor)
        return new

    def revise(self, mapping_id: str, *, actor: str, target_statement: str,
               statement_section: str, note_number: str, subnote_code: str | None) -> MappingDecision:
        old = self._decisions.get(mapping_id)
        if old is None:
            raise MappingWorkflowError("unknown mapping_id")
        if not actor.strip() or not target_statement or not statement_section or not note_number:
            raise MappingWorkflowError("actor and complete target are required")
        new = replace(old, target_statement=target_statement,
                      statement_section=statement_section, note_number=note_number,
                      subnote_code=subnote_code, status="Draft", reviewer=None,
                      review_date=None,
                      version=old.version + 1)
        self._record(old, new, actor)
        return new

    def carry_forward(self, previous: Iterable[MappingDecision], *, actor: str) -> int:
        if not actor.strip():
            raise MappingWorkflowError("actor required")
        candidates = tuple(previous)
        ids = [d.mapping_id for d in candidates]
        if len(ids) != len(set(ids)) or any(i in self._decisions for i in ids):
            raise MappingWorkflowError("duplicate carry-forward mapping_id")
        for d in candidates:
            if d.company_code != self.company_code or d.fiscal_year != self.fiscal_year - 1:
                raise MappingWorkflowError("carry-forward source company/year mismatch")
            if d.status != "Approved":
                raise MappingWorkflowError("only approved mappings may be carried forward")
        for d in candidates:
            self.propose(replace(d, fiscal_year=self.fiscal_year, status="Draft",
                                 reviewer=None, review_date=None, version=1,
                                 source="CarryForward"), actor=actor)
        return len(candidates)


def import_mapping_workbook(content: bytes, company_code: str, fiscal_year: int) -> tuple[MappingDecision, ...]:
    """Read only preflight-valid mapping rows; reject any out-of-scope rows."""
    report = preflight_pilot_workbook(content)
    if not report.valid:
        raise MappingWorkflowError(f"workbook preflight failed: {report.issues[0].code}")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        columns = SHEET_HEADERS["Mapping"]
        decisions: list[MappingDecision] = []
        for values in workbook["Mapping"].iter_rows(min_row=2, values_only=True):
            if not any(v is not None for v in values):
                continue
            row = dict(zip(columns, values))
            if (row["company_code"], row["fiscal_year"]) != (company_code, fiscal_year):
                raise MappingWorkflowError("import includes foreign company/year")
            decisions.append(MappingDecision(
                row["mapping_id"], row["company_code"], row["fiscal_year"],
                row["account_full_code"], row["target_statement"],
                row["statement_section"], row["note_number"], row["subnote_code"],
                row["mapping_version"], row["approval_status"], row["reviewer"],
                row["source"], row["review_date"],
            ))
        return tuple(decisions)
    finally:
        workbook.close()


def export_mapping_rows(decisions: Iterable[MappingDecision]) -> bytes:
    """Export mapping decisions to a standalone review sheet (not an approved Pilot input)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mapping"
    columns = ("mapping_id", "company_code", "fiscal_year", "account_full_code",
               "target_statement", "statement_section", "note_number", "subnote_code",
               "mapping_version", "approval_status", "reviewer", "source")
    sheet.append(columns)
    for decision in decisions:
        sheet.append((decision.mapping_id, decision.company_code, decision.fiscal_year,
                      decision.account_full_code, decision.target_statement,
                      decision.statement_section, decision.note_number, decision.subnote_code,
                      decision.version, decision.status, decision.reviewer, decision.source))
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@dataclass(frozen=True)
class UnknownAccount:
    company_code: str
    fiscal_year: int
    account_full_code: str


def unknown_account_queue(
    accounts: Iterable[tuple[str, int, str]],
    decisions: Iterable[MappingDecision],
    *,
    company_code: str,
    fiscal_year: int,
) -> tuple[UnknownAccount, ...]:
    """Return unapproved/unmapped accounts, deduplicated and deterministically sorted.

    A Draft or Reviewed mapping does not remove an account from the queue.
    Foreign-scope account input fails closed rather than silently disappearing.
    """
    if not company_code.strip() or isinstance(fiscal_year, bool) or fiscal_year <= 0:
        raise MappingWorkflowError("invalid queue company/year")
    approved: set[str] = set()
    for decision in decisions:
        if (decision.company_code, decision.fiscal_year) != (company_code, fiscal_year):
            raise MappingWorkflowError("foreign mapping in unknown-account queue")
        if decision.status == "Approved":
            approved.add(decision.account_full_code)
    unknown: set[str] = set()
    for company, year, account in accounts:
        if (company, year) != (company_code, fiscal_year):
            raise MappingWorkflowError("foreign account in unknown-account queue")
        if not isinstance(account, str) or not account.strip():
            raise MappingWorkflowError("invalid account code in unknown-account queue")
        if account not in approved:
            unknown.add(account)
    return tuple(UnknownAccount(company_code, fiscal_year, account)
                 for account in sorted(unknown))


def export_pilot_mapping_workbook(
    source: bytes, decisions: Iterable[MappingDecision], *,
    company_code: str, fiscal_year: int,
) -> bytes:
    """Produce a complete Pilot workbook, validated for re-import."""
    check = preflight_pilot_workbook(source)
    if not check.valid:
        raise MappingWorkflowError("source workbook failed preflight")
    workbook = load_workbook(BytesIO(source))
    try:
        headers = SHEET_HEADERS["Mapping"]
        mapping_sheet = workbook["Mapping"]
        accounts = {}
        for values in workbook["Accounts"].iter_rows(min_row=2, values_only=True):
            if any(v is not None for v in values):
                row = dict(zip(SHEET_HEADERS["Accounts"], values))
                accounts[(row["company_code"], row["fiscal_year"], row["account_full_code"])] = row
        notes = set()
        for values in workbook["Notes"].iter_rows(min_row=2, values_only=True):
            if any(v is not None for v in values):
                row = dict(zip(SHEET_HEADERS["Notes"], values))
                notes.add((row["company_code"], row["fiscal_year"], row["statement"],
                           row["statement_section"], row["note_number"], row["subnote_code"]))
        previous = {}
        foreign_rows = []
        for values in mapping_sheet.iter_rows(min_row=2, values_only=True):
            if any(v is not None for v in values):
                row = dict(zip(headers, values))
                previous[row["mapping_id"]] = row
                if (row["company_code"], row["fiscal_year"]) != (company_code, fiscal_year):
                    foreign_rows.append(tuple(row[field] for field in headers))
        replacements = []
        seen = set()
        for d in decisions:
            if (d.company_code, d.fiscal_year) != (company_code, fiscal_year):
                raise MappingWorkflowError("foreign company/year in export")
            if d.mapping_id in seen:
                raise MappingWorkflowError("duplicate mapping ID in export")
            seen.add(d.mapping_id)
            account = accounts.get((company_code, fiscal_year, d.account_full_code))
            if account is None:
                raise MappingWorkflowError("unknown account in export")
            if (company_code, fiscal_year, d.target_statement, d.statement_section,
                d.note_number, d.subnote_code) not in notes:
                raise MappingWorkflowError("unknown note/subnote in export")
            if d.status not in {"Draft", "Reviewed", "Approved"}:
                raise MappingWorkflowError("invalid mapping status")
            if d.status != "Draft" and (not d.reviewer or not d.review_date):
                raise MappingWorkflowError("review evidence required")
            row = {field: None for field in headers}
            row.update({
                "mapping_id": d.mapping_id, "company_code": company_code,
                "fiscal_year": fiscal_year, "account_full_code": d.account_full_code,
                "account_level": previous[d.mapping_id]["account_level"]
                if d.mapping_id in previous else "General",
                "target_statement": d.target_statement,
                "statement_section": d.statement_section,
                "note_number": d.note_number, "subnote_code": d.subnote_code,
                "mapping_version": d.version, "approval_status": d.status,
                "source": d.source, "reviewer": d.reviewer,
                "review_date": d.review_date,
                "validation_status": "Valid" if d.status == "Approved" else "Pending",
            })
            for field in ("general_code", "subsidiary_code", "detail1_code",
                          "detail2_code", "detail3_code", "detail4_code"):
                row[field] = account[field]
            replacements.append(tuple(row[field] for field in headers))
        if mapping_sheet.max_row > 1:
            mapping_sheet.delete_rows(2, mapping_sheet.max_row - 1)
        for values in foreign_rows + replacements:
            mapping_sheet.append(values)
        stream = BytesIO()
        workbook.save(stream)
        output = stream.getvalue()
    finally:
        workbook.close()
    validation = preflight_pilot_workbook(output)
    if not validation.valid:
        raise MappingWorkflowError(f"export preflight failed: {validation.issues[0].code}")
    return output
