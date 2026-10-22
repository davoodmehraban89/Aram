"""Strict, non-executing XLSX interchange for explicit mapping decisions.

The workbook carries only mapping identifiers, not balances or inferred targets.
All database changes are staged in a savepoint; the caller owns the outer commit.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from aram_backend.mapping_service import MappingValidationError, save_mapping_rule

HEADERS = ("rule_key", "account_code", "target_note_id", "mapping_type", "version")
SHEET = "Mapping"
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_ROWS = 10_000


@dataclass(frozen=True)
class MappingEntry:
    rule_key: str
    account_code: str
    target_note_id: int
    mapping_type: str
    version: int


def _safe_archive(content: bytes) -> None:
    if not isinstance(content, bytes) or not content or len(content) > MAX_FILE_BYTES:
        raise MappingValidationError("XLSX must be nonempty and at most 10 MiB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 1_000 or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
                raise MappingValidationError("XLSX archive exceeds safety limits")
            if any(i.flag_bits & 0x1 for i in infos):
                raise MappingValidationError("Encrypted XLSX entries are not supported")
    except (BadZipFile, ValueError) as exc:
        raise MappingValidationError("Invalid XLSX archive") from exc


def _parse_int(value: object, label: str, row_number: int) -> int:
    if type(value) is not int or value <= 0:
        raise MappingValidationError(f"Row {row_number}: {label} must be a positive integer")
    return value


def _parse_text(value: object, label: str, limit: int, row_number: int) -> str:
    if not isinstance(value, str) or not value or value.strip() != value or len(value) > limit:
        raise MappingValidationError(f"Row {row_number}: invalid {label}")
    return value


def preview_mapping_workbook(content: bytes) -> tuple[MappingEntry, ...]:
    """Read and validate an XLSX mapping sheet without modifying any database."""
    _safe_archive(content)
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (OSError, ValueError, KeyError, BadZipFile) as exc:
        raise MappingValidationError("Unreadable XLSX workbook") from exc
    try:
        if book.sheetnames != [SHEET]:
            raise MappingValidationError("Workbook must contain only a Mapping sheet")
        sheet = book[SHEET]
        if sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1:
            raise MappingValidationError("Too many mapping rows")
        if sheet.max_column is not None and sheet.max_column > len(HEADERS):
            raise MappingValidationError("Unexpected columns in Mapping sheet")
        rows = sheet.iter_rows(values_only=False)
        header = next(rows, ())
        if tuple(cell.value for cell in header) != HEADERS:
            raise MappingValidationError("Mapping sheet headers do not match the contract")
        entries: list[MappingEntry] = []
        seen: set[tuple[str, int]] = set()
        for row_number, cells in enumerate(rows, start=2):
            if row_number > MAX_ROWS + 1:
                raise MappingValidationError("Too many mapping rows")
            if any(cell.data_type == "f" for cell in cells):
                raise MappingValidationError(f"Row {row_number}: formulas are prohibited")
            values = tuple(cell.value for cell in cells)
            if all(value is None for value in values):
                continue
            if len(values) != len(HEADERS):
                raise MappingValidationError(f"Row {row_number}: invalid column count")
            entry = MappingEntry(
                rule_key=_parse_text(values[0], "rule_key", 128, row_number),
                account_code=_parse_text(values[1], "account_code", 64, row_number),
                target_note_id=_parse_int(values[2], "target_note_id", row_number),
                mapping_type=_parse_text(values[3], "mapping_type", 32, row_number),
                version=_parse_int(values[4], "version", row_number),
            )
            identity = (entry.rule_key, entry.version)
            if identity in seen:
                raise MappingValidationError(f"Row {row_number}: duplicate rule key/version")
            seen.add(identity)
            entries.append(entry)
        if not entries:
            raise MappingValidationError("Mapping sheet contains no mapping entries")
        return tuple(entries)
    finally:
        book.close()


def export_mapping_workbook(entries: tuple[MappingEntry, ...]) -> bytes:
    """Create a deterministic-structure XLSX; force text cells to prevent formula execution."""
    if not entries or len(entries) > MAX_ROWS:
        raise MappingValidationError("Mapping export must contain 1 to 10000 entries")
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(HEADERS)
    seen: set[tuple[str, int]] = set()
    for index, entry in enumerate(entries, start=2):
        if not isinstance(entry, MappingEntry):
            raise MappingValidationError(f"Row {index}: invalid mapping entry")
        key = (_parse_text(entry.rule_key, "rule_key", 128, index),
               _parse_int(entry.version, "version", index))
        if key in seen:
            raise MappingValidationError(f"Row {index}: duplicate rule key/version")
        seen.add(key)
        row = (
            key[0], _parse_text(entry.account_code, "account_code", 64, index),
            _parse_int(entry.target_note_id, "target_note_id", index),
            _parse_text(entry.mapping_type, "mapping_type", 32, index), key[1],
        )
        sheet.append(row)
        for col in (1, 2, 4):
            sheet.cell(index, col).data_type = "s"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def import_mapping_workbook(session: Session, *, year_id: int, content: bytes) -> int:
    """Stage all valid mappings atomically, without committing the caller's transaction."""
    entries = preview_mapping_workbook(content)
    with session.begin_nested():
        for entry in entries:
            save_mapping_rule(
                session, year_id=year_id, rule_key=entry.rule_key,
                account_code=entry.account_code, target_note_id=entry.target_note_id,
                mapping_type=entry.mapping_type, version=entry.version,
            )
    return len(entries)
