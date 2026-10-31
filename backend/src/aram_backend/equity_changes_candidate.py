"""Unsigned statement of changes in equity candidate.

Read-only, exact-Decimal arithmetic over a scoped Pilot workbook. Requires
explicit note-level component, movement classification and debit-net sign
policies. It does not infer the nature of an equity transaction or issue an
official financial statement.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
from typing import Any, Sequence

from openpyxl import load_workbook

from aram_backend.coverage_preview import (
    CoverageError, _assert_visible, _guard_archive, preview_mapping_coverage,
)
from aram_backend.pilot_workbook import SHEET_HEADERS
from aram_backend.reconciliation_review import (
    NoteReference, _money, _sum_exact, reconcile_workbook_for_review,
)

MAX_POLICIES = 10_000
EQUITY_COMPONENTS = frozenset({
    "SHARE_CAPITAL", "SHARE_PREMIUM", "RETAINED_EARNINGS",
    "REVALUATION_SURPLUS", "OTHER_RESERVES", "NONCONTROLLING_INTEREST",
    "OTHER_EQUITY",
})
MOVEMENT_CATEGORIES = frozenset({
    "CAPITAL_CONTRIBUTION", "NET_PROFIT_LOSS", "DIVIDENDS",
    "OTHER_COMPREHENSIVE_INCOME", "TRANSFER",
    "PRIOR_PERIOD_ADJUSTMENT", "OTHER_MOVEMENT",
})


@dataclass(frozen=True)
class EquityMovementPolicy:
    statement_section: str
    note_number: str
    subnote_code: str
    equity_component: str
    movement_category: str
    sign_multiplier: int
    line_label: str
    policy_evidence_ref: str


@dataclass(frozen=True)
class EquityAccountMovement:
    account_code: str
    account_name: str
    source_excel_row: int
    source_row: int
    statement_section: str
    note_number: str
    subnote_code: str
    equity_component: str
    movement_category: str
    sign_multiplier: int
    opening_balance: str
    debit_turnover: str
    credit_turnover: str
    movement_amount: str
    closing_balance: str
    policy_evidence_ref: str


@dataclass(frozen=True)
class EquityNoteMovement:
    statement_section: str
    note_number: str
    subnote_code: str
    equity_component: str
    movement_category: str
    line_label: str
    opening_balance: str
    movement_amount: str
    closing_balance: str
    source_closing_exposure: str
    note_reference_evidence: str
    policy_evidence_ref: str


@dataclass(frozen=True)
class EquityComponentTotal:
    equity_component: str
    opening_balance: str
    movement_amount: str
    closing_balance: str


@dataclass(frozen=True)
class EquityChangesCandidate:
    company_code: str
    fiscal_year: int
    source_sha256: str
    note_references_sha256: str
    presentation_policies_sha256: str
    opening_equity_total: str
    capital_contributions: str
    net_profit_loss_movements: str
    dividends: str
    other_comprehensive_income: str
    transfers: str
    prior_period_adjustments: str
    other_movements: str
    total_movement: str
    closing_equity_total: str
    rollforward_difference: str
    component_totals: tuple[EquityComponentTotal, ...]
    note_movements: tuple[EquityNoteMovement, ...]
    account_movements: tuple[EquityAccountMovement, ...]
    findings: tuple[str, ...]
    eligible_for_analyst_review: bool
    kind: str = "UNAPPROVED_EQUITY_CHANGES_PRESENTATION_CANDIDATE"
    official_financial_statement: bool = False
    accounting_policy_authorized: bool = False
    source_ledger_completeness_verified: bool = False
    independently_reviewed: bool = False
    export_authorized: bool = False
    releasable: bool = False

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _text(value: object, label: str, *, blank_ok: bool = False) -> str:
    if (not isinstance(value, str) or len(value) > 256
            or value != value.strip() or (not value and not blank_ok)
            or any(ord(char) < 32 or ord(char) == 127 for char in value)):
        raise CoverageError("INVALID_EQUITY_POLICY_" + label)
    return value


def _policy_key(
    policy: EquityMovementPolicy,
) -> tuple[str, str, str]:
    return policy.statement_section, policy.note_number, policy.subnote_code


def _validate_policies(
    policies: Sequence[EquityMovementPolicy],
) -> tuple[dict[tuple[str, str, str], EquityMovementPolicy], str]:
    if not isinstance(policies, (tuple, list)) or len(policies) > MAX_POLICIES:
        raise CoverageError("INVALID_EQUITY_POLICIES")
    selected: dict[tuple[str, str, str], EquityMovementPolicy] = {}
    for policy in policies:
        if not isinstance(policy, EquityMovementPolicy):
            raise CoverageError("INVALID_EQUITY_POLICY")
        _text(policy.statement_section, "SECTION")
        _text(policy.note_number, "NOTE")
        _text(policy.subnote_code, "SUBNOTE", blank_ok=True)
        _text(policy.line_label, "LABEL")
        _text(policy.policy_evidence_ref, "EVIDENCE")
        if (not isinstance(policy.equity_component, str)
                or policy.equity_component not in EQUITY_COMPONENTS):
            raise CoverageError("INVALID_EQUITY_COMPONENT")
        if (not isinstance(policy.movement_category, str)
                or policy.movement_category not in MOVEMENT_CATEGORIES):
            raise CoverageError("INVALID_EQUITY_MOVEMENT_CATEGORY")
        if type(policy.sign_multiplier) is not int or policy.sign_multiplier not in (-1, 1):
            raise CoverageError("INVALID_EQUITY_SIGN")
        key = _policy_key(policy)
        if key in selected:
            raise CoverageError("DUPLICATE_EQUITY_POLICY")
        selected[key] = policy
    canonical = [asdict(selected[key]) for key in sorted(selected)]
    digest = sha256(json.dumps(
        canonical, ensure_ascii=True, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    return selected, digest


def build_equity_changes_candidate(
    workbook_bytes: bytes, *, company_code: str, fiscal_year: int,
    references: Sequence[NoteReference],
    policies: Sequence[EquityMovementPolicy],
) -> EquityChangesCandidate:
    """Create a nonofficial equity movement candidate from explicit note policies.

    Source balances are assumed to be debit-net signed. Each note's sign
    multiplier converts its opening, debit-credit movement and closing to
    presentation sign. All accounts assigned to one note inherit that note's
    single movement classification; use separate subnotes when a note mixes
    transaction types. That constraint is intentionally explicit.
    """
    selected, policy_hash = _validate_policies(policies)
    preview = preview_mapping_coverage(
        workbook_bytes, company_code=company_code, fiscal_year=fiscal_year,
    )
    review = reconcile_workbook_for_review(
        workbook_bytes, company_code=company_code,
        fiscal_year=fiscal_year, references=references,
    )
    if (review.kind != "UNAPPROVED_NOTE_RECONCILIATION_REVIEW"
            or review.financial_statement or review.trusted_approvals
            or review.approved or review.releasable
            or preview.financial_statement or preview.trusted_approvals
            or preview.releasable
            or preview.source_sha256 != review.source_sha256):
        raise CoverageError("EQUITY_REVIEW_PROVENANCE_INVALID")

    eq_notes = {
        (note.section, note.note_number, note.subnote_code): note
        for note in preview.notes if note.statement == "EQ"
    }
    findings: list[str] = []
    if not review.matched_for_analyst_review:
        findings.append("UPSTREAM_RECONCILIATION_BLOCKED")
    if not eq_notes:
        findings.append("NO_EQUITY_NOTES")
    if set(eq_notes) - set(selected):
        findings.append("MISSING_EQUITY_POLICY")
    if set(selected) - set(eq_notes):
        findings.append("ORPHAN_EQUITY_POLICY")

    ref_by_key = {
        (note.statement, note.section, note.note_number, note.subnote_code): note
        for note in review.notes
    }
    trace_by_note: dict[tuple[str, str, str], list[Any]] = {}
    for trace in preview.mappings:
        if trace.included and trace.target_statement == "EQ":
            key = (trace.target_section or "", trace.target_note or "",
                   trace.target_subnote or "")
            trace_by_note.setdefault(key, []).append(trace)

    try:
        book = load_workbook(
            BytesIO(workbook_bytes), read_only=False,
            data_only=False, keep_links=False,
        )
    except Exception as exc:
        raise CoverageError("EQUITY_WORKBOOK_INVALID") from exc
    try:
        for sheet in book.worksheets:
            _assert_visible(sheet)
        for name, headers in SHEET_HEADERS.items():
            if name not in book or tuple(cell.value for cell in book[name][1]) != headers:
                raise CoverageError("EQUITY_WORKBOOK_INVALID")

        account_rows: dict[str, tuple[int, dict[str, Any]]] = {}
        sheet = book["Accounts"]
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), 2,
        ):
            row = dict(zip(SHEET_HEADERS["Accounts"], values))
            if not any(value is not None and str(value).strip() for value in row.values()):
                continue
            if (isinstance(row.get("company_code"), str)
                    and row["company_code"].strip() == review.company_code
                    and type(row.get("fiscal_year")) is int
                    and row["fiscal_year"] == review.fiscal_year):
                code = row.get("account_full_code")
                if not isinstance(code, str) or not code.strip():
                    raise CoverageError("EQUITY_ACCOUNT_CODE_INVALID")
                code = code.strip()
                if code in account_rows:
                    raise CoverageError("EQUITY_DUPLICATE_ACCOUNT")
                account_rows[code] = (row_number, row)

        note_lines: list[EquityNoteMovement] = []
        account_lines: list[EquityAccountMovement] = []
        component_opening: dict[str, list[Decimal]] = {
            item: [] for item in sorted(EQUITY_COMPONENTS)
        }
        component_movement: dict[str, list[Decimal]] = {
            item: [] for item in sorted(EQUITY_COMPONENTS)
        }
        component_closing: dict[str, list[Decimal]] = {
            item: [] for item in sorted(EQUITY_COMPONENTS)
        }
        movement_buckets: dict[str, list[Decimal]] = {
            item: [] for item in sorted(MOVEMENT_CATEGORIES)
        }
        opening_values: list[Decimal] = []
        closing_values: list[Decimal] = []
        movement_values: list[Decimal] = []

        for key in sorted(set(eq_notes) & set(selected)):
            note = eq_notes[key]
            policy = selected[key]
            matched = ref_by_key.get(("EQ", *key))
            if matched is None or not matched.matched:
                findings.append("EQUITY_NOTE_REFERENCE_NOT_MATCHED")
                continue
            traces = trace_by_note.get(key, [])
            if not traces:
                findings.append("EQUITY_NOTE_WITHOUT_ACCOUNTS")
                continue
            note_openings: list[Decimal] = []
            note_debits: list[Decimal] = []
            note_credits: list[Decimal] = []
            note_closings: list[Decimal] = []
            for trace in traces:
                source = account_rows.get(trace.account_code)
                if source is None:
                    raise CoverageError("EQUITY_SOURCE_ACCOUNT_MISSING")
                excel_row, row = source
                if excel_row != trace.account_excel_row:
                    raise CoverageError("EQUITY_SOURCE_ROW_MISMATCH")
                opening, debit, credit, closing = (
                    _money(str(row[field]))
                    for field in ("opening", "debit", "credit", "closing")
                )
                source_movement = _sum_exact([debit, -credit])
                if _sum_exact([opening, source_movement, -closing]) != 0:
                    raise CoverageError("EQUITY_ACCOUNT_ROLLFORWARD_INVALID")
                if debit < 0 or credit < 0:
                    raise CoverageError("EQUITY_NEGATIVE_TURNOVER")
                sign = policy.sign_multiplier
                shown_opening = opening * sign
                shown_movement = source_movement * sign
                shown_closing = closing * sign
                note_openings.append(opening)
                note_debits.append(debit)
                note_credits.append(credit)
                note_closings.append(closing)
                component_opening[policy.equity_component].append(shown_opening)
                component_movement[policy.equity_component].append(shown_movement)
                component_closing[policy.equity_component].append(shown_closing)
                movement_buckets[policy.movement_category].append(shown_movement)
                account_lines.append(EquityAccountMovement(
                    account_code=trace.account_code,
                    account_name=trace.account_name,
                    source_excel_row=excel_row,
                    source_row=trace.source_row,
                    statement_section=key[0], note_number=key[1],
                    subnote_code=key[2],
                    equity_component=policy.equity_component,
                    movement_category=policy.movement_category,
                    sign_multiplier=sign,
                    opening_balance=str(shown_opening),
                    debit_turnover=str(debit),
                    credit_turnover=str(credit),
                    movement_amount=str(shown_movement),
                    closing_balance=str(shown_closing),
                    policy_evidence_ref=policy.policy_evidence_ref,
                ))
            opening_raw = _sum_exact(note_openings)
            debit_raw = _sum_exact(note_debits)
            credit_raw = _sum_exact(note_credits)
            closing_raw = _sum_exact(note_closings)
            if _money(note.closing_exposure) != closing_raw:
                raise CoverageError("EQUITY_NOTE_CLOSING_MISMATCH")
            opening_shown = opening_raw * policy.sign_multiplier
            movement_shown = _sum_exact([debit_raw, -credit_raw]) * policy.sign_multiplier
            closing_shown = closing_raw * policy.sign_multiplier
            if _sum_exact([opening_shown, movement_shown, -closing_shown]) != 0:
                findings.append("EQUITY_NOTE_ROLLFORWARD_MISMATCH")
            opening_values.append(opening_shown)
            movement_values.append(movement_shown)
            closing_values.append(closing_shown)
            note_lines.append(EquityNoteMovement(
                statement_section=key[0], note_number=key[1],
                subnote_code=key[2], equity_component=policy.equity_component,
                movement_category=policy.movement_category,
                line_label=policy.line_label,
                opening_balance=str(opening_shown),
                movement_amount=str(movement_shown),
                closing_balance=str(closing_shown),
                source_closing_exposure=str(closing_raw),
                note_reference_evidence=matched.evidence_ref or "",
                policy_evidence_ref=policy.policy_evidence_ref,
            ))

        opening_total = _sum_exact(opening_values)
        movement_total = _sum_exact(movement_values)
        closing_total = _sum_exact(closing_values)
        rollforward = _sum_exact([opening_total, movement_total, -closing_total])
        if rollforward != 0:
            findings.append("EQUITY_TOTAL_ROLLFORWARD_MISMATCH")

        component_totals = tuple(
            EquityComponentTotal(
                equity_component=component,
                opening_balance=str(_sum_exact(component_opening[component])),
                movement_amount=str(_sum_exact(component_movement[component])),
                closing_balance=str(_sum_exact(component_closing[component])),
            )
            for component in sorted(EQUITY_COMPONENTS)
            if (component_opening[component] or component_movement[component]
                or component_closing[component])
        )
        movement_totals = {
            category: _sum_exact(values)
            for category, values in movement_buckets.items()
        }
        if not any(note_lines):
            findings.append("NO_EQUITY_MOVEMENTS")
        return EquityChangesCandidate(
            company_code=review.company_code, fiscal_year=review.fiscal_year,
            source_sha256=review.source_sha256,
            note_references_sha256=review.reference_sha256,
            presentation_policies_sha256=policy_hash,
            opening_equity_total=str(opening_total),
            capital_contributions=str(movement_totals["CAPITAL_CONTRIBUTION"]),
            net_profit_loss_movements=str(movement_totals["NET_PROFIT_LOSS"]),
            dividends=str(movement_totals["DIVIDENDS"]),
            other_comprehensive_income=str(movement_totals["OTHER_COMPREHENSIVE_INCOME"]),
            transfers=str(movement_totals["TRANSFER"]),
            prior_period_adjustments=str(movement_totals["PRIOR_PERIOD_ADJUSTMENT"]),
            other_movements=str(movement_totals["OTHER_MOVEMENT"]),
            total_movement=str(movement_total),
            closing_equity_total=str(closing_total),
            rollforward_difference=str(rollforward),
            component_totals=component_totals,
            note_movements=tuple(note_lines),
            account_movements=tuple(sorted(
                account_lines, key=lambda item: item.account_code,
            )),
            findings=tuple(sorted(set(findings))),
            eligible_for_analyst_review=not findings and bool(note_lines),
        )
    finally:
        book.close()
