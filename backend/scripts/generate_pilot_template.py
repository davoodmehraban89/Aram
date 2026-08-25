from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from aram_backend.pilot_workbook import SHEET_HEADERS


OUTPUT = Path(__file__).resolve().parents[2] / "samples" / "Aram_Pilot_Input_and_Mapping_Template_v0.1.xlsx"


def build_template(output: Path = OUTPUT) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for name, headers in SHEET_HEADERS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).coordinate}"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for column in range(1, len(headers) + 1):
            sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 18

    guide = workbook["Guide_FA"]
    instructions = [
        ("شروع", "ابتدا اطلاعات شرکت و سال مالی را در برگه Company وارد کنید."),
        ("کدها", "همه کدها را به‌صورت متن وارد کنید تا صفرهای ابتدایی حفظ شوند."),
        ("مبالغ", "مبالغ را فقط از منبع واقعی وارد کنید؛ فایل نمونه فاقد مبلغ ساختگی است."),
        ("کنترل", "قبل از ورود اطلاعات، Preflight را اجرا و همه خطاهای اعلام‌شده را رفع کنید."),
        ("گردش", "وضعیت نگاشت فقط Draft، Reviewed یا Approved است و خودکار ارتقا نمی‌یابد."),
    ]
    for row in instructions:
        guide.append(row)
    guide.sheet_view.rightToLeft = True

    lists = workbook["Lists"]
    approved_lists = {
        "company_status": ("Draft", "Approved", "Locked"),
        "review_status": ("Draft", "Reviewed", "Approved"),
        "validation_status": ("Valid", "Invalid", "Pending"),
        "test_status": ("Pass", "Fail", "Blocked", "Not Run"),
        "severity": ("Critical", "High", "Medium", "Low"),
    }
    for list_name, values in approved_lists.items():
        for value in values:
            lists.append((list_name, value))

    checks = workbook["Checks"]
    for row in (
        ("REQUIRED_FIELDS", "Required-field completeness", "PENDING", 0, "Run backend preflight"),
        ("DUPLICATE_ACCOUNTS", "Duplicate accounts", "PENDING", 0, "Run backend preflight"),
        ("UNKNOWN_MAPPING_ACCOUNT", "Unknown mapping accounts", "PENDING", 0, "Run backend preflight"),
        ("INVALID_MAPPING", "Invalid mappings", "PENDING", 0, "Run backend preflight"),
        ("NON_ZERO_CLOSING_DELTA", "Non-zero closing_delta", "PENDING", 0, "Run backend preflight"),
    ):
        checks.append(row)

    company_status = DataValidation(type="list", formula1='"Draft,Approved,Locked"')
    review_status = DataValidation(type="list", formula1='"Draft,Reviewed,Approved"')
    workbook["Company"].add_data_validation(company_status)
    company_status.add("G2:G500")
    workbook["Accounts"].add_data_validation(review_status)
    review_status.add("V2:V5000")
    mapping_status = DataValidation(type="list", formula1='"Draft,Reviewed,Approved"')
    workbook["Mapping"].add_data_validation(mapping_status)
    mapping_status.add("Q2:Q5000")
    red = PatternFill("solid", fgColor="FFC7CE")
    workbook["Accounts"].conditional_formatting.add(
        "W2:W5000", FormulaRule(formula=["W2<>0"], fill=red)
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


if __name__ == "__main__":
    print(build_template())
